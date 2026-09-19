import json
import hashlib
import importlib
import os
import zipfile
from pathlib import Path

import openpyxl
import pytest

from app.main import main
import app.offline_review as offline_review
from app.offline_review import review_directory, to_json, to_markdown, write_outputs


def _write_bid(path: Path, total: int, *, identical_text: bool = True, unit: str = "m³",
               specification: str = "一般土方") -> None:
    path.mkdir(parents=True, exist_ok=True)
    (path / "报价.csv").write_text(
        "项目名称,规格,单位,数量,单价,合价\n"
        f"土方,{specification},{unit},1,30000,30000\n"
        f"混凝土,{specification},{unit},2,20000,40000\n"
        f"钢筋,{specification},{unit},1,30000,30000\n"
        f"投标报价,{total}\n",
        encoding="utf-8",
    )
    text = "施工组织设计统一说明。\n" * 20 if identical_text else "不同施工方案说明。\n" * 20
    (path / "说明.txt").write_text(text, encoding="utf-8")


def test_review_directory_excludes_relation_file_and_emits_risk_signals(tmp_path):
    _write_bid(tmp_path / "甲公司", 100000)
    _write_bid(tmp_path / "乙公司", 100300)
    _write_bid(tmp_path / "丙公司", 120000, identical_text=False)
    relation = tmp_path / "relations.csv"
    relation.write_text(
        "bidder_a,bidder_b,relation,source\n甲公司,乙公司,法定代表人相同,本地依法取得的登记资料\n",
        encoding="utf-8",
    )

    result = review_directory(tmp_path, project="测试项目", relations=relation)

    assert {x["name"] for x in result["bidders"]} == {"甲公司", "乙公司", "丙公司"}
    assert result["summary"]["file_count"] == 6
    assert result["relation_source"]["sha256"]
    codes = {signal["code"] for signal in result["signals"]}
    assert {"QUOTE_NEAR_MATCH", "QUOTE_OUTLIER", "SYNCHRONIZED_LINE_ITEMS", "LOCAL_RELATION_CLUE"} <= codes
    assert all(signal["auto_conclusion"] is False for signal in result["signals"])
    assert result["summary"]["manual_review_required"] is True


def test_relation_json_respects_depth_budget(tmp_path, monkeypatch):
    relation = tmp_path / "relations.json"
    relation.write_text(
        json.dumps({"relations": [{"bidder_a": "甲", "bidder_b": "乙", "relation": "相同"}]}),
        encoding="utf-8",
    )
    monkeypatch.setattr(offline_review, "MAX_JSON_DEPTH", 2)

    result = review_directory(tmp_path, relations=relation)

    assert result["summary"]["parse_error_count"] == 1
    assert "JSON 解析受限" in result["parse_errors"][0]["error"]


def test_review_directory_integrates_new_pattern_signals_and_control_quote(tmp_path):
    item_rows = [
        "项目名称,规格,单位,数量,单价,合价",
        *(f"清单项{i},一般土方,m³,1,100,100" for i in range(1, 9)),
    ]
    for name, total in (("甲公司", 900000), ("乙公司", 900100), ("丙公司", 900200)):
        bidder = tmp_path / name
        bidder.mkdir()
        (bidder / "报价.csv").write_text(
            "\n".join(item_rows + [
                f"投标报价,{total}",
                "招标控制价,1000000",
                "项目经理,张三",
                "联系电话,13800000000",
            ]) + "\n",
            encoding="utf-8",
        )
    relations = tmp_path / "relations.csv"
    relations.write_text(
        "bidder_a,bidder_b,relation,source\n"
        "甲公司,乙公司,两家法定代表人系亲兄弟,本地登记资料\n",
        encoding="utf-8",
    )

    result = review_directory(tmp_path, relations=relations)
    codes = {signal["code"] for signal in result["signals"]}
    expected = {
        "QUOTE_ARITHMETIC_PATTERN", "UNIFORM_DISCOUNT_RATE", "LINE_ITEM_SET_MATCH",
        "PERSON_OVERLAP", "KINSHIP_RELATION",
    }

    assert expected <= codes
    assert all(
        signal["auto_conclusion"] is False
        for signal in result["signals"]
        if signal["code"] in expected
    )
    assert all(signal["legal_basis"] for signal in result["signals"])
    assert result["summary"]["manual_review_required"] is True

    bidders = {bidder["name"]: bidder for bidder in result["bidders"]}
    for bidder in bidders.values():
        assert bidder["primary_quote"]["kind"] == "explicit_total"
        assert bidder["primary_quote"]["label"] == "投标报价"
        control_quotes = [quote for quote in bidder["quotes"] if quote["kind"] == "control"]
        assert len(control_quotes) == 1
        assert control_quotes[0]["label"] == "招标控制价"
        assert bidder["primary_quote"]["kind"] != control_quotes[0]["kind"]


def test_payment_account_is_not_exposed_in_public_result(tmp_path):
    account = "6222021234567890123"
    for name, total in (("甲公司", 100), ("乙公司", 200)):
        bidder = tmp_path / name
        bidder.mkdir()
        (bidder / "报价.csv").write_text(
            f"银行账号,{account}\n投标报价,{total}\n", encoding="utf-8"
        )

    result = review_directory(tmp_path)
    payload = to_json(result)

    assert "PAYMENT_ACCOUNT_MATCH" in payload
    assert account not in payload
    assert all("bank_account" not in bidder["metadata"] for bidder in result["bidders"])


def test_csv_quote_uses_first_amount_before_tax_rate(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.csv").write_text("投标报价,1000,税率,13\n", encoding="utf-8")
    ambiguous = tmp_path / "乙"
    ambiguous.mkdir()
    (ambiguous / "报价.csv").write_text("bid_price,tax_rate,13,100000\n", encoding="utf-8")

    result = review_directory(tmp_path)

    bidders = {bidder["name"]: bidder for bidder in result["bidders"]}
    assert bidders["甲"]["primary_quote"]["value"] == 1000
    assert bidders["乙"]["primary_quote"] is None


def test_csv_quote_skips_adjacent_unlabeled_amounts(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.csv").write_text("bid_price,13,100000\n", encoding="utf-8")

    result = review_directory(tmp_path)

    assert result["bidders"][0]["primary_quote"] is None
    assert result["bidders"][0]["quotes"] == []


def test_text_quote_skips_multiple_amounts_in_one_line(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.txt").write_text("投标报价：税率13%，金额100000\n", encoding="utf-8")

    result = review_directory(tmp_path)

    assert result["bidders"][0]["primary_quote"] is None
    assert not result["bidders"][0]["quotes"]


def test_text_quote_prefers_currency_amount_before_label_over_trailing_duration(tmp_path):
    """真实标书句式「（¥ 金额）的投标总报价，工期 N 日历天」：金额在标签之前。"""
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "商务文件.txt").write_text(
        "1.我方已仔细研究了某工程施工招标文件的全部内容，愿意以人民币（大写）"
        " 柒仟玖佰柒拾玖万陆仟伍佰伍拾玖元壹角捌分 元（¥ 79796559.18 ）的投标总报价，"
        "工期 1124 日历天，按合同约定实施和完成承包工程。\n",
        encoding="utf-8",
    )

    result = review_directory(tmp_path)

    quotes = result["bidders"][0]["quotes"]
    assert len(quotes) == 1
    assert quotes[0]["value"] == 79796559.18
    assert quotes[0]["label"] == "总报价"
    assert result["bidders"][0]["primary_quote"]["value"] == 79796559.18


def test_text_quote_skips_duration_amount_without_currency_prefix(tmp_path):
    """标签后唯一金额紧邻工期词语且无前置货币金额时，宁可漏报不当作报价。"""
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.txt").write_text("总报价，工期 1124 日历天\n", encoding="utf-8")

    result = review_directory(tmp_path)

    assert not result["bidders"][0]["quotes"]
    assert result["bidders"][0]["primary_quote"] is None


def test_review_parses_json_and_xlsx_with_traceable_hashes(tmp_path):
    first = tmp_path / "甲"
    first.mkdir()
    (first / "报价.json").write_text(
        json.dumps({"投标总价": "1.2万元", "items": [{"项目名称": "土方", "合价": 1200}]}),
        encoding="utf-8",
    )
    second = tmp_path / "乙"
    second.mkdir()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "报价表"
    sheet.append(["项目名称", "数量", "单价", "合价"])
    sheet.append(["土方", 1, 1200, 1200])
    sheet.append(["混凝土", 1, 800, 800])
    sheet.append(["钢筋", 1, 1000, 1000])
    sheet.append(["投标报价", 3000, None, None])
    sheet.row_dimensions[2].hidden = True
    sheet["F1"] = "=SUM(B2:B4)"
    workbook.properties.creator = "测试作者"
    workbook.save(second / "报价.xlsx")

    result = review_directory(tmp_path)
    bidders = {x["name"]: x for x in result["bidders"]}
    assert bidders["甲"]["primary_quote"]["value"] == 12000
    assert bidders["乙"]["primary_quote"]["value"] == 3000
    assert bidders["乙"]["files"][0]["sha256"]
    assert "author" in bidders["乙"]["metadata"]
    assert bidders["乙"]["files"][0]["structure"]["sheet_count"] == 1
    sheet_meta = bidders["乙"]["files"][0]["structure"]["sheets"][0]
    assert sheet_meta["hidden_row_count"] == 1
    assert sheet_meta["formula_count"] == 1
    assert bidders["乙"]["line_items"][0]["comparability_status"] == "INSUFFICIENT_DATA"
    json.loads(to_json(result))


def test_xlsx_formula_without_cached_value_is_explicit_warning(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["投标报价", "=100000"])
    workbook.save(bidder / "报价.xlsx")

    result = review_directory(tmp_path)
    bidder_result = result["bidders"][0]
    sheet_meta = bidder_result["files"][0]["structure"]["sheets"][0]

    assert bidder_result["primary_quote"] is None
    assert sheet_meta["formula_count"] == 1
    assert any("不计算公式" in warning["error"] for warning in result["parse_warnings"])
    assert "解析提示" in to_markdown(result)


def test_xlsx_native_numeric_values_keep_decimal_scale(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["项目名称", "单位", "规格", "数量", "单价", "合价"])
    sheet.append(["混凝土", "m3", "C30", 1.234, 1000, 1234])
    sheet.append(["投标报价", 123.456])
    workbook.save(bidder / "报价.xlsx")

    result = review_directory(tmp_path)
    bidder_result = result["bidders"][0]

    assert bidder_result["primary_quote"]["value"] == 123.456
    assert bidder_result["line_items"][0]["quantity"] == 1.234
    assert bidder_result["line_items"][0]["amount"] == 1234


def test_note_lines_never_produce_quote_or_control_evidence(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.txt").write_text(
        "注：本表适用于建设项目招标控制价或投标报价的汇总。第1页共2页\n"
        "注：1.“名称、规格、型号”、“基本价格指数”栏由招标人填写，价格指数 22\n"
        "投标总价: 100000\n",
        encoding="utf-8",
    )
    result = review_directory(tmp_path)
    bidder_info = result["bidders"][0]
    values = [q["value"] for q in bidder_info["quotes"]]
    assert values == [100000]
    assert bidder_info["primary_quote"]["value"] == 100000


def test_pdf_text_layer_is_parsed_and_blank_pdf_flagged(tmp_path, monkeypatch):
    monkeypatch.setattr(offline_review, "_ocr_pdf_page", lambda *args: "")
    monkeypatch.setattr(offline_review, "_pdf_ocr_tools", lambda: ("pdftoppm", "tesseract", "chi_sim+eng", {}))
    import io

    from pypdf import PdfWriter
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    # 默认 Helvetica 渲染不了中文（会变■），注册 CID 中文字体贴近真实标书。
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except Exception:
        pass

    def _build_pdf(path, lines):
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer)
        try:
            c.setFont("STSong-Light", 12)
        except Exception:
            pass
        y = 780
        for line in lines:
            c.drawString(72, y, line)
            y -= 24
        c.save()
        path.write_bytes(buffer.getvalue())

    bidder = tmp_path / "甲公司"
    bidder.mkdir()
    # 只画一行：drawString 的行在 extract_text 里可能粘连成一行，多数字会被报价识别保守拒绝。
    _build_pdf(bidder / "经济标.pdf", ["投标总价: 185986265.34 yuan"])

    result = review_directory(tmp_path)
    info = result["bidders"][0]["files"][0]
    assert info["parse_status"] == "OK"
    assert info["structure"]["page_count"] == 1
    assert info["structure"]["text_truncated"] is False
    values = [q["value"] for q in result["bidders"][0]["quotes"]]
    assert 185986265.34 in values

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    blank = tmp_path / "乙公司"
    blank.mkdir()
    blank_buffer = io.BytesIO()
    writer.write(blank_buffer)
    (blank / "扫描件.pdf").write_bytes(blank_buffer.getvalue())
    result2 = review_directory(tmp_path)
    scanned = next(f for b in result2["bidders"] for f in b["files"]
                   if f["path"].endswith("扫描件.pdf"))
    assert scanned["parse_status"] == "PARTIAL"
    assert any("文本层" in w["reason"] for w in scanned.get("parse_warnings", []))


def test_pdf_producer_device_metadata_triggers_e01(tmp_path):
    """同一物理设备（Producer/Creator）产出两家投标 PDF 是典型围标线索，
    PDF 文档信息字典须进入 E-01 的电子元数据比对池。"""
    import io

    from pypdf import PdfWriter

    for company, producer in (("甲公司", "RICOH MP C2003"), ("乙公司", "RICOH MP C2003"),
                              ("丙公司", "Fuji Xerox D110")):
        directory = tmp_path / company
        directory.mkdir()
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        writer.add_metadata({"/Producer": producer, "/Creator": f"{producer} driver"})
        buffer = io.BytesIO()
        writer.write(buffer)
        (directory / "投标文件.pdf").write_bytes(buffer.getvalue())

    result = review_directory(tmp_path)

    device_signals = [s for s in result["signals"]
                      if s["code"] == "METADATA_MATCH"
                      and any(v["field"] == "device" for v in s["evidence"])]
    scopes = {frozenset(s["scope"].split(" ↔ ")) for s in device_signals}
    assert frozenset({"甲公司", "乙公司"}) in scopes
    assert not any("丙公司" in scope for scope in scopes)


def test_quote_extraction_review_findings_regression(tmp_path):
    """独立复核发现的取值优先级与抗损边界（offline_review._text_quotes）。

    - tail 唯一合法金额优先：head 里的更早 ¥ 金额（如单价）不得顶替合计；
    - head 货币金额取最靠近标签的一个（最右），且其与标签间不得夹限定词；
    - OCR 断号金额（¥ 797 965 59.18）不猜测；
    - head 货币金额吸收「万元/亿」后缀，与 tail 口径一致；
    - 工期语境窗口 ±16 字符。
    """
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.txt").write_text(
        "单价：¥ 350.00 合计金额 ¥ 87,500.00\n"
        "投标保证金：¥ 50000，愿以（¥ 79796559.18）的投标总报价\n"
        "（¥ 797 965 59.18 ）的投标总报价，工期 100 天\n"
        "（¥ 7979.66 万元）的投标总报价\n"
        "总报价见商务标，工期要求：自开工之日起算 1124 天内完工\n",
        encoding="utf-8",
    )

    result = review_directory(tmp_path)

    values = sorted(q["value"] for q in result["bidders"][0]["quotes"])
    # 87500（合计）、79796559.18（最靠近标签的货币金额）、79796600（万元换算）
    assert values == [87500.0, 79796559.18, 79796600.0]
    # 断号与工期语境的两行不得产出报价


def test_quote_owner_noun_amounts_never_become_quotes(tmp_path):
    """回归（限定词/属主复核）：保证金、单价等属主名词紧邻的金额不是报价。

    - 标签后唯一金额前紧邻属主名词（「投标保证金 50000 元」「单价 350 元」）
      时不得充当报价——此前 tail 侧无限定词守卫，50000 会被当作总报价；
    - head 内最靠近标签的货币金额前紧邻属主名词（「保证金（¥ 5000）」）时
      回溯更早金额，全部属主化则宁漏勿错；
    - 属主名词与金额距离较远（「履约保证金另行提交，本报价 50000」）或
      修饰标签本身的词（「含税」）不阻断正常取值。"""
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.txt").write_text(
        "总报价见商务标，投标保证金 50000.00 元\n"
        "投标总报价：单价 350.00 元\n"
        "报价人民币（¥ 79796559.18），其中保证金（¥ 5000.00 元）的投标总报价\n"
        "投标总报价：履约保证金另行提交，本报价 50000.00 元\n"
        "投标报价（含税）：50000 元\n",
        encoding="utf-8",
    )

    result = review_directory(tmp_path)

    values = sorted(q["value"] for q in result["bidders"][0]["quotes"])
    assert values == [50000.0, 50000.0]


def test_table_paths_skip_owner_nouns_and_split_amounts():
    """表格路径同样受属主名词与断号金额守卫约束（offline_review._rows_quotes）。"""
    quotes, _, _, _ = offline_review._rows_quotes(
        [["投标总报价（单价：350 元）"], ["投标总价", "50000"]], "表.xlsx")
    assert [q["value"] for q in quotes] == [50000.0]

    quotes, _, _, _ = offline_review._rows_quotes([["投标总价", "50000. 18"]], "表.xlsx")
    assert quotes == []


def test_ocr_split_amount_boundaries(tmp_path):
    """回归（OCR 截断边界复核）：小数点/千分位与数字被空白断开的金额不猜。

    「¥ 79796559. 18」修复前在 head 侧被解析为 79796559.0、tail 侧解析为
    18.0；「1, 234.56」在 tail 侧解析为 234.56；JSON 字段「50000. 18」
    解析为 18。修复后 txt head/tail 与 JSON 字段一律宁漏勿错；
    千分位无空格的「1,234,567.89」不受影响。"""
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.txt").write_text(
        "（¥ 79796559. 18）的投标总报价\n"
        "投标总报价：¥ 79796559. 18\n"
        "投标总报价：1, 234.56\n"
        "投标总价：1,234,567.89 元\n",
        encoding="utf-8",
    )
    (bidder / "报价.json").write_text(
        json.dumps({"投标总价": "50000. 18"}), encoding="utf-8")

    result = review_directory(tmp_path)

    values = [q["value"] for q in result["bidders"][0]["quotes"]]
    assert values == [1234567.89]
    assert result["bidders"][0]["primary_quote"]["value"] == 1234567.89


def test_head_currency_amount_unit_suffixes(tmp_path):
    """回归（金额单位复核）：head 货币金额吸收「亿元/万/元」后缀，
    与标签后金额（tail）及 _parse_amount 口径一致。"""
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.txt").write_text(
        "（¥ 0.8 亿元）的投标总报价\n"
        "（¥ 100 万）的投标总报价\n"
        "（¥ 50000 元）的投标总报价\n",
        encoding="utf-8",
    )

    result = review_directory(tmp_path)

    values = sorted(q["value"] for q in result["bidders"][0]["quotes"])
    assert values == [50000.0, 1000000.0, 80000000.0]


def test_duration_context_window_boundary(tmp_path):
    """回归（工期窗口边界复核）：工期词语在唯一金额 ±16 字符窗口内不取值
    （test_quote_extraction_review_findings_regression 的 1124 天内完工行），
    窗口外不得误伤正常报价，与
    test_normal_single_amount_unaffected_by_duration_guard 合成双向语义。"""
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.txt").write_text(
        "投标总价：79796559.18 元，上述报价包含暂列金额与专业工程暂估价，工期另见专用条款\n",
        encoding="utf-8",
    )

    result = review_directory(tmp_path)

    assert [q["value"] for q in result["bidders"][0]["quotes"]] == [79796559.18]


def test_public_result_redacts_accounts_inside_raw_fields(tmp_path):
    account = "6222021234567890123"
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.csv").write_text(
        f"投标报价,100000,银行账号,{account}\n", encoding="utf-8")

    result = review_directory(tmp_path)
    assert account not in to_json(result)

    json_bidder = tmp_path / "乙"
    json_bidder.mkdir()
    (json_bidder / "清单.json").write_text(
        json.dumps({"项目名称": "混凝土", "单位": "m3", "规格": "C30",
                    "合价": 1234, "bank_account": account}, ensure_ascii=False),
        encoding="utf-8",
    )
    result = review_directory(tmp_path)
    assert account not in to_json(result)


def test_line_items_do_not_match_by_name_when_units_differ(tmp_path):
    _write_bid(tmp_path / "甲", 100000, unit="m³")
    _write_bid(tmp_path / "乙", 100300, unit="吨")

    result = review_directory(tmp_path)

    assert not any(signal["code"] == "SYNCHRONIZED_LINE_ITEMS" for signal in result["signals"])
    assert result["summary"]["comparable_line_item_count"] == 6
    assert all(
        item["comparison_key"]
        for bidder in result["bidders"]
        for item in bidder["line_items"]
    )


def test_line_items_do_not_match_by_name_when_specifications_differ(tmp_path):
    _write_bid(tmp_path / "甲", 100000, specification="C30")
    _write_bid(tmp_path / "乙", 100300, specification="C40")

    result = review_directory(tmp_path)

    assert not any(signal["code"] == "SYNCHRONIZED_LINE_ITEMS" for signal in result["signals"])


def test_line_items_with_missing_unit_or_specification_are_marked_not_comparable(tmp_path):
    for name in ("甲", "乙"):
        bidder = tmp_path / name
        bidder.mkdir()
        (bidder / "报价.csv").write_text(
            "项目名称,数量,单价,合价\n"
            "土方,1,30000,30000\n"
            "混凝土,2,20000,40000\n"
            "钢筋,1,30000,30000\n",
            encoding="utf-8",
        )

    result = review_directory(tmp_path)
    items = [item for bidder in result["bidders"] for item in bidder["line_items"]]

    assert len(items) == 6
    assert all(item["comparability_status"] == "INSUFFICIENT_DATA" for item in items)
    assert all(item["comparison_key"] is None for item in items)
    assert result["summary"]["insufficient_line_item_count"] == 6
    assert result["summary"]["line_item_comparability"]["auto_merge"] is False
    assert not any(signal["code"] == "SYNCHRONIZED_LINE_ITEMS" for signal in result["signals"])


def test_json_composite_total_is_not_parsed_as_scalar_quote(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.json").write_text(
        json.dumps({
            "bid_price": {"amount": 12345, "currency": "CNY"},
            "total": {"value": "12345元", "currency": "CNY"},
            "items": [{"项目名称": "土方", "单位": "m³", "规格": "一般土方", "合价": 100}],
        }, ensure_ascii=False),
        encoding="utf-8",
    )

    result = review_directory(tmp_path)
    bidder_result = result["bidders"][0]

    assert bidder_result["primary_quote"] is None
    assert bidder_result["quotes"] == []
    assert bidder_result["line_items"][0]["comparability_status"] == "COMPARABLE"


def test_quote_parser_keeps_invalid_currency_and_negative_values_traceable(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.csv").write_text(
        "项目名称,规格,单位,合价\n"
        "土方,一般土方,m³,100\n"
        "投标报价,-5000\n"
        "报价金额,100 USD\n",
        encoding="utf-8",
    )

    result = review_directory(tmp_path)

    assert result["bidders"][0]["primary_quote"] is None
    assert result["bidders"][0]["quotes"] == []
    warnings = result["parse_warnings"]
    assert any("负值" in warning["error"] for warning in warnings)
    assert any("非人民币" in warning["error"] for warning in warnings)
    assert result["summary"]["parse_warning_count"] == 2


def test_normalized_text_hash_covers_content_after_compare_prefix(tmp_path):
    prefix = "a" * (offline_review.MAX_TEXT_FOR_COMPARE + 100)
    for name, suffix in (("甲", "left-tail"), ("乙", "right-tail")):
        bidder = tmp_path / name
        bidder.mkdir()
        (bidder / "说明.txt").write_text(prefix + suffix, encoding="utf-8")

    result = review_directory(tmp_path)
    files = [bidder["files"][0] for bidder in result["bidders"]]

    assert files[0]["normalized_sha256"] != files[1]["normalized_sha256"]
    assert not any(signal["code"] == "TEXT_EXACT_MATCH" for signal in result["signals"])


def test_short_identical_files_do_not_create_exact_match_signal(tmp_path):
    for name in ("甲", "乙"):
        bidder = tmp_path / name
        bidder.mkdir()
        (bidder / "空.txt").write_text("", encoding="utf-8")

    result = review_directory(tmp_path)

    assert not any(signal["code"] == "TEXT_EXACT_MATCH" for signal in result["signals"])


def test_json_item_prefers_total_over_unit_price(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "报价.json").write_text(
        json.dumps({"items": [{"项目名称": "土方", "单位": "m³", "规格": "一般土方",
                                "单价": 12, "合价": 120}]}),
        encoding="utf-8",
    )

    result = review_directory(tmp_path)

    assert result["bidders"][0]["line_items"][0]["amount"] == 120


def test_scan_skips_and_records_symlink(tmp_path):
    target = tmp_path / "outside.txt"
    target.write_text("投标报价,100", encoding="utf-8")
    bidder = tmp_path / "甲"
    bidder.mkdir()
    link = bidder / "链接.txt"
    try:
        link.symlink_to(target)
    except (OSError, NotImplementedError) as exc:
        pytest.skip(f"当前环境不允许创建符号链接：{exc}")

    result = review_directory(tmp_path)

    assert not any(file["path"].endswith("链接.txt") for bidder in result["bidders"] for file in bidder["files"])
    assert any(item["path"].endswith("链接.txt") and item["reason"] == "SYMLINK_SKIPPED"
               for item in result["skipped_files"])
    assert result["summary"]["skipped_file_count"] == 1
    assert result["scan"]["symlink_count"] == 1


def test_scan_stops_at_file_budget(tmp_path, monkeypatch):
    monkeypatch.setattr(offline_review, "MAX_INPUT_FILES", 1)
    (tmp_path / "a.txt").write_text("投标报价,100", encoding="utf-8")
    (tmp_path / "b.txt").write_text("投标报价,200", encoding="utf-8")

    result = review_directory(tmp_path)

    assert result["scan"]["truncated"] is True
    assert any(item["kind"] == "file_count" for item in result["scan"]["budget_errors"])
    assert result["summary"]["scan_error_count"] == 1
    assert result["summary"]["file_count"] == 1


def test_scan_stops_at_byte_budget(tmp_path, monkeypatch):
    monkeypatch.setattr(offline_review, "MAX_INPUT_BYTES", 10)
    (tmp_path / "a.txt").write_bytes(b"123456")
    (tmp_path / "b.txt").write_bytes(b"123456")

    result = review_directory(tmp_path)

    assert result["scan"]["truncated"] is True
    assert any(item["kind"] == "bytes" for item in result["scan"]["budget_errors"])
    assert result["summary"]["file_count"] == 1


def test_xlsx_zip_guard_runs_before_openpyxl(tmp_path, monkeypatch):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    archive_path = bidder / "恶意.xlsx"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("xl/large.bin", b"123456")
    monkeypatch.setattr(offline_review, "MAX_XLSX_ENTRY_UNCOMPRESSED_BYTES", 5)

    result = review_directory(tmp_path)

    assert result["summary"]["parse_error_count"] == 1
    assert "XLSX 压缩包成员" in result["parse_errors"][0]["error"]


def test_xlsx_read_pass_respects_audit_cell_budget(tmp_path, monkeypatch):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["项目名称", "合价"])
    sheet.append(["土方", 100])
    workbook.save(bidder / "报价.xlsx")
    monkeypatch.setattr(offline_review, "MAX_XLSX_AUDIT_CELLS", 2)

    result = review_directory(tmp_path)
    sheet_meta = result["bidders"][0]["files"][0]["structure"]["sheets"][0]

    assert sheet_meta["read_truncated"] is True
    assert sheet_meta["audit_truncated"] is True
    assert sheet_meta["formula_count"] is None


def test_unsupported_and_broken_files_are_not_silently_clean(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    (bidder / "附件.xls").write_bytes(b"legacy workbook")
    (bidder / "坏.json").write_text("{not-json", encoding="utf-8")

    result = review_directory(tmp_path)

    assert result["summary"]["unsupported_file_count"] == 1
    assert result["summary"]["parse_error_count"] == 1
    assert any(".xls" in x["path"] for x in result["unsupported_files"])
    assert any(x["path"].endswith("坏.json") for x in result["parse_errors"])


def test_broken_json_parse_error_keeps_file_fingerprint(tmp_path):
    bidder = tmp_path / "甲"
    bidder.mkdir()
    broken = bidder / "坏.json"
    broken.write_text("{not-json", encoding="utf-8")
    raw = broken.read_bytes()

    result = review_directory(tmp_path)
    file_meta = result["bidders"][0]["files"][0]
    parse_error = result["parse_errors"][0]

    expected_sha256 = hashlib.sha256(raw).hexdigest()
    assert file_meta["size_bytes"] == len(raw)
    assert file_meta["sha256"] == expected_sha256
    assert parse_error["size_bytes"] == len(raw)
    assert parse_error["sha256"] == expected_sha256


def test_markdown_and_json_keep_manual_only_boundary(tmp_path):
    _write_bid(tmp_path / "甲", 100000)
    result = review_directory(tmp_path)

    report = to_markdown(result)
    payload = json.loads(to_json(result))
    assert "不构成串通投标、违法" in report
    assert payload["legal_notice"]["conclusion"].startswith("仅输出风险预警")
    assert payload["summary"]["manual_review_required"] is True


def test_cli_review_bids_writes_json_and_markdown(tmp_path, capsys):
    _write_bid(tmp_path / "甲", 100000)
    output_dir = tmp_path.parent / f"{tmp_path.name}-out"
    json_path = output_dir / "review.json"
    report_path = output_dir / "review.md"

    rc = main([
        "review-bids", str(tmp_path), "--output-json", str(json_path),
        "--output-report", str(report_path),
    ])

    assert rc == 0
    assert json_path.is_file() and report_path.is_file()
    assert "风险信号" in capsys.readouterr().out


def test_write_outputs_rejects_input_paths_before_writing(tmp_path, capsys):
    _write_bid(tmp_path / "甲", 100000)
    result = review_directory(tmp_path)
    source_path = tmp_path / "甲" / "报价.csv"
    before = source_path.read_bytes()
    new_output_dir = tmp_path / "new-output"

    with pytest.raises(ValueError, match="输入目录"):
        write_outputs(result, json_path=tmp_path)

    with pytest.raises(ValueError, match="输入目录"):
        write_outputs(
            result,
            json_path=source_path,
            report_path=new_output_dir / "report.md",
        )

    assert source_path.read_bytes() == before
    assert not new_output_dir.exists()

    rc = main([
        "review-bids", str(tmp_path), "--output-json", str(source_path),
    ])
    cli_output = capsys.readouterr().out
    assert rc == 2
    assert "Traceback" not in cli_output
    assert "无法完成离线审查" in cli_output
    assert source_path.read_bytes() == before


def test_write_outputs_rejects_hardlink_alias(tmp_path):
    _write_bid(tmp_path / "甲", 100000)
    result = review_directory(tmp_path)
    source_path = tmp_path / "甲" / "报价.csv"
    alias = tmp_path.parent / f"{tmp_path.name}-alias.csv"
    try:
        os.link(source_path, alias)
    except OSError as exc:
        pytest.skip(f"当前环境不允许创建硬链接：{exc}")

    before = source_path.read_bytes()
    with pytest.raises(ValueError, match="输入源文件"):
        write_outputs(result, json_path=alias)
    assert source_path.read_bytes() == before


def test_web_offline_review_entry_reads_only_explicit_directory(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient
    import app.web.server as server

    _write_bid(tmp_path / "甲", 100000)
    monkeypatch.setenv("BQC_REVIEW_ROOT", str(tmp_path))
    importlib.reload(server)
    client = TestClient(server.app)

    page = client.get("/review-bids")
    result = client.post("/review-bids", data={"input_dir": str(tmp_path), "project": "Web 测试"})

    assert page.status_code == 200
    assert "离线审查" in page.text
    assert result.status_code == 200
    assert "人工复核" in result.text
    assert "Web 测试" in result.text

    outside = tmp_path.parent / f"{tmp_path.name}-outside"
    _write_bid(outside / "乙", 100000)
    rejected = client.post("/review-bids", data={"input_dir": str(outside)})
    assert rejected.status_code == 400
    assert "BQC_REVIEW_ROOT" in rejected.text
    assert rejected.headers["content-type"].startswith("text/html")


def test_web_offline_review_accepts_subdirectory_of_review_root(tmp_path, monkeypatch):
    """回归：审查根的“子目录”是正常使用形态。

    背景（2026-09-18 真实界面验收）：限根逻辑写成了
    ``inside = path == root if allow_root else root in path.parents``，
    运算符优先级使 allow_root=True 时只放行根目录本身，任何子目录都被 400，
    /review-bids 页面自加限制起核心流程不可用。锁定语义：根或其子目录均放行。
    """
    from fastapi.testclient import TestClient
    import app.web.server as server

    proj = tmp_path / "某市政项目"
    _write_bid(proj / "甲", 100000)
    _write_bid(proj / "乙", 100300)
    monkeypatch.setenv("BQC_REVIEW_ROOT", str(tmp_path))
    importlib.reload(server)
    client = TestClient(server.app)

    result = client.post("/review-bids", data={"input_dir": str(proj), "project": "子目录审查"})
    assert result.status_code == 200
    assert "人工复核" in result.text
    assert "子目录审查" in result.text


def test_web_offline_review_error_renders_html_not_bare_json(tmp_path, monkeypatch):
    """表单页校验失败必须回到人可读的 HTML 页面（含原因与可修正的表单），不得裸输出 JSON。"""
    from fastapi.testclient import TestClient
    import app.web.server as server

    monkeypatch.setenv("BQC_REVIEW_ROOT", str(tmp_path))
    importlib.reload(server)
    client = TestClient(server.app)

    outside = tmp_path.parent / f"{tmp_path.name}-out2"
    r = client.post("/review-bids", data={"input_dir": str(outside)})
    assert r.status_code == 400
    assert r.headers["content-type"].startswith("text/html")
    assert "BQC_REVIEW_ROOT" in r.text  # 服务端给出的原因原样保留
    assert "开始离线审查" in r.text  # 表单仍在，用户可直接修正重试


def test_root_files_without_explicit_prefix_stay_one_bidder(tmp_path):
    """回归（2026-09-20 合成探针）：同一投标人的多个无前缀根文件不得按文件名拆成
    多家“投标人”互相比较——那会凭空制造「报价 ↔ 施工组织」式假 F-01 信号。

    语义按 review_directory 文档锁定：无「投标人__文件名」前缀的根文件一律归入
    根目录名对应的单一投标人。"""
    (tmp_path / "报价.txt").write_text("投标总价：1000000\n", encoding="utf-8")
    (tmp_path / "施工组织.md").write_text(
        "# 施工组织设计\n投标总价：1000000\n第一节 总体部署\n" + "内容" * 60 + "\n",
        encoding="utf-8",
    )

    result = review_directory(tmp_path)

    assert [b["name"] for b in result["bidders"]] == [tmp_path.name]
    assert len(result["bidders"][0]["files"]) == 2
    assert not [s for s in result["signals"] if s["code"].startswith("F-")]


def test_root_files_with_explicit_prefix_still_use_prefix(tmp_path):
    """「投标人__文件名.ext」前缀约定保持不变（README 布局向后兼容）。"""
    (tmp_path / "甲__报价.csv").write_text("投标总价,1000000\n", encoding="utf-8")
    (tmp_path / "乙__报价.csv").write_text("投标总价,1001000\n", encoding="utf-8")

    result = review_directory(tmp_path)

    assert {b["name"] for b in result["bidders"]} == {"甲", "乙"}
    assert any(s["code"] == "QUOTE_NEAR_MATCH" for s in result["signals"])


def test_vertical_summary_pairs_label_line_with_next_amount_line(tmp_path):
    """回归：纵排汇总（标签独占一行、金额在下一非空行）此前产不出任何报价。

    常见于导出的单列报价汇总。配对必须可追溯：locator 指向金额所在行，
    证据原文同时保留标签行与金额行。"""
    (tmp_path / "甲").mkdir()
    (tmp_path / "乙").mkdir()
    (tmp_path / "甲" / "报价.txt").write_text(
        "投标报价汇总表\n\n投标总价\n1000000 元\n", encoding="utf-8")
    (tmp_path / "乙" / "报价.txt").write_text(
        "投标报价汇总表\n\n投标总价\n1001000 元\n", encoding="utf-8")

    result = review_directory(tmp_path)
    quotes = {b["name"]: b["quotes"] for b in result["bidders"]}
    assert [q["value"] for q in quotes["甲"]] == [1000000.0]
    assert quotes["甲"][0]["locator"] == "第4行"
    assert "第3行:投标总价" in quotes["甲"][0]["raw"]
    assert any(s["code"] == "QUOTE_NEAR_MATCH" for s in result["signals"])


def test_transposed_summary_pairs_labels_with_value_row_by_position(tmp_path):
    """回归：转置汇总（标签行+数值行）此前产不出任何报价。

    配对须按位置一一对应；控制价只记 kind=control，不得成为主报价，
    也不参与 F-01（控制价两家本来就相同）。"""
    import io

    import openpyxl

    for name, values in (("甲", (500, 1000000)), ("乙", (500, 1001000))):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["招标控制价", "投标总价"])
        sheet.append(list(values))
        buffer = io.BytesIO()
        workbook.save(buffer)
        (tmp_path / name).mkdir()
        (tmp_path / name / "报价.xlsx").write_bytes(buffer.getvalue())

    result = review_directory(tmp_path)
    for bidder in result["bidders"]:
        kinds = {q["label"]: q["kind"] for q in bidder["quotes"]}
        assert kinds == {"招标控制价": "control", "投标总价": "explicit_total"}
        assert bidder["primary_quote"]["value"] in (1000000.0, 1001000.0)
    assert any(s["code"] == "QUOTE_NEAR_MATCH" for s in result["signals"])


def test_vertical_pairing_skips_when_amounts_continue_as_a_column(tmp_path):
    """回归：数值行之后仍紧跟数字行时更像一列明细数字，不得把列首数字当总价。

    源自合成探针 Q8：公式单元格被读取层丢弃后，标签行下方留下明细数字列，
    曾被配成「投标总价=500」的假报价。守卫生效时不产报价，公式提示保留。"""
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["投标总价", "=SUM(B2:B3)"])
    sheet.append([None, 500])
    sheet.append([None, 999500])
    buffer = io.BytesIO()
    workbook.save(buffer)
    (tmp_path / "甲").mkdir()
    (tmp_path / "甲" / "报价.xlsx").write_bytes(buffer.getvalue())

    result = review_directory(tmp_path)

    assert result["bidders"][0]["quotes"] == []
    assert result["bidders"][0]["primary_quote"] is None
    assert any("公式" in w["error"] for w in result["parse_warnings"])


def test_vertical_pairing_stays_silent_on_ambiguous_followup_line(tmp_path):
    """标签行下一行带税率/百分比语义或多个数字时一律放弃配对，宁可漏报。"""
    (tmp_path / "甲").mkdir()
    (tmp_path / "甲" / "报价.txt").write_text(
        "投标总价\n下浮率 5%\n", encoding="utf-8")

    result = review_directory(tmp_path)

    assert result["bidders"][0]["quotes"] == []


def test_relation_file_ambiguities_are_traceable_not_silent(tmp_path):
    """回归：关联线索文件的格式歧义必须留痕，不得静默吞掉。

    - CSV 表头重复列名：dict(zip()) 只保留最后一个取值，历史行为保留但须提示；
    - JSON 非对象行：此前被静默忽略，现提示忽略行数。"""
    _write_bid(tmp_path / "甲公司", 100000)
    _write_bid(tmp_path / "乙公司", 100300)

    dup = tmp_path / "重复表头.csv"
    dup.write_text(
        "bidder_a,bidder_a,relation\n甲公司,乙公司,法定代表人相同\n", encoding="utf-8")
    result_dup = review_directory(tmp_path, relations=dup)
    assert any("重复列名" in w["error"] for w in result_dup["parse_warnings"])
    assert any(s["code"] == "LOCAL_RELATION_CLUE" for s in result_dup["signals"])

    mixed = tmp_path / "混合行.json"
    mixed.write_text(
        json.dumps(["甲公司", {"bidder_a": "甲公司", "bidder_b": "乙公司", "relation": "同址"}],
                   ensure_ascii=False),
        encoding="utf-8")
    result_mixed = review_directory(tmp_path, relations=mixed)
    assert any("非对象记录" in w["error"] for w in result_mixed["parse_warnings"])
    assert len(result_mixed["relation_clues"]) == 1
    assert result_mixed["summary"]["parse_warning_count"] >= 1


def test_normal_single_amount_unaffected_by_duration_guard(tmp_path):
    """回归（工期守卫复核）：正常唯一金额报价不受工期守卫影响。

    守卫结构（offline_review._text_quotes）中 quotes.append 与 value/label
    赋值同处 `if not _DURATION_CONTEXT_RE.search(...)` 分支内；本测试与
    test_text_quote_skips_duration_amount_without_currency_prefix 合起来
    锁死「工期数字跳过、正常金额照常识别」的双向语义。"""
    (tmp_path / "甲").mkdir()
    (tmp_path / "甲" / "报价.txt").write_text("投标总价：1000000 元\n", encoding="utf-8")

    result = review_directory(tmp_path)

    assert [q["value"] for q in result["bidders"][0]["quotes"]] == [1000000.0]
    assert result["bidders"][0]["primary_quote"]["value"] == 1000000.0


def test_duration_context_guard_also_applies_to_adjacent_cells(tmp_path):
    """回归：CSV 相邻单元格带工期语义（完工天等）时不得产报价，
    与 txt 路径 _DURATION_CONTEXT_RE 的守卫范围保持一致。"""
    (tmp_path / "甲").mkdir()
    (tmp_path / "甲" / "q.csv").write_text("投标总价,2000 完工天\n", encoding="utf-8")

    result = review_directory(tmp_path)

    assert result["bidders"][0]["quotes"] == []


def test_gbk_encoded_quote_file_is_not_silently_mojibake(tmp_path):
    """回归：GBK 编码（大陆遗留系统常见）的报价文件不得被静默解成乱码。

    偶数长度且不含未配对代理的字节流会「成功」通过无 BOM 的 'utf-16' 解码器
    （按本机字节序），整文件变乱码且不留任何解析提示——报价/联系人标签全部
    丢失，parse_status 仍为 OK。UTF-16 只认显式 BOM，GBK 交由 gb18030 兜底。"""
    raw = "联系人：张三，电话 13800000000\n投标报价：50000 元\n".encode("gb18030")
    # 前置条件：这段字节流恰好落入无 BOM 'utf-16' 的「成功乱码」陷阱。
    with pytest.raises(UnicodeDecodeError):
        raw.decode("utf-8-sig")
    raw.decode("utf-16")  # 无 BOM 也能解出（乱码），证明陷阱真实存在
    assert not raw.startswith((b"\xff\xfe", b"\xfe\xff"))
    (tmp_path / "甲公司").mkdir()
    (tmp_path / "甲公司" / "报价.txt").write_bytes(raw)

    result = review_directory(tmp_path)

    bidder = result["bidders"][0]
    assert bidder["files"][0]["parse_status"] == "OK"
    assert [q["value"] for q in bidder["quotes"]] == [50000.0]
    assert bidder["primary_quote"]["value"] == 50000.0
    assert not result["parse_errors"]
    assert not result["parse_warnings"]


def test_utf16_bom_quote_file_still_parses(tmp_path):
    """带 BOM 的 UTF-16 报价文件继续正常解析（既有能力不回归）。"""
    (tmp_path / "甲公司").mkdir()
    (tmp_path / "甲公司" / "报价.txt").write_text(
        "联系人：李四\n投标报价：60000 元\n", encoding="utf-16")

    result = review_directory(tmp_path)

    bidder = result["bidders"][0]
    assert bidder["files"][0]["parse_status"] == "OK"
    assert [q["value"] for q in bidder["quotes"]] == [60000.0]
    assert bidder["primary_quote"]["value"] == 60000.0
