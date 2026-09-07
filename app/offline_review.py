"""离线多投标文件风险预警。

本模块只读本地 txt/md/csv/json/xlsx，输出可追溯的客观相似性信号。
它不访问网络、不修改输入文件、不判定串通投标或投标无效；所有信号均须人工复核。
"""
from __future__ import annotations

import csv
import difflib
import hashlib
import io
import json
import math
import os
import re
import stat
import zipfile
from datetime import datetime, timezone
from itertools import islice
from pathlib import Path
from statistics import median, mean, pstdev
from typing import Any, Iterable

from .offline_review_patterns import RULE_IDS as PATTERN_RULE_IDS, apply as apply_pattern_signals


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".xlsx"}
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
MAX_TEXT_FOR_COMPARE = 120_000
MAX_INPUT_FILES = 500
MAX_INPUT_BYTES = 200 * 1024 * 1024
MAX_XLSX_ZIP_MEMBERS = 2_000
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_XLSX_ENTRY_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_JSON_DEPTH = 100
MAX_JSON_NODES = 100_000
MAX_XLSX_AUDIT_CELLS = 200_000

# 阈值是预警筛选参数，不是法律推定。调整时应同步更新 README/验收测试。
THRESHOLDS = {
    "near_quote_relative_diff": 0.005,
    "low_dispersion_cv": 0.01,
    "outlier_relative_diff": 0.10,
    "line_item_match_relative_diff": 0.005,
    "text_similarity": 0.92,
    "structure_similarity": 0.80,
    "min_common_line_items": 3,
    "min_text_chars": 80,
}

RULE_IDS = {
    "QUOTE_NEAR_MATCH": "F-01",
    "QUOTE_LOW_DISPERSION": "F-02",
    "QUOTE_OUTLIER": "F-03",
    "SYNCHRONIZED_LINE_ITEMS": "F-04",
    "METADATA_MATCH": "E-01",
    "TEXT_EXACT_MATCH": "S-01",
    "TEXT_HIGH_SIMILARITY": "S-02",
    "STRUCTURE_SIMILARITY": "S-03",
    "LOCAL_RELATION_CLUE": "P-01",
}
RULE_IDS.update(PATTERN_RULE_IDS)

_ELECTRONIC_METADATA_FIELDS = {"author", "machine_id", "mac", "ip", "disk_serial", "certificate"}

_AMOUNT_RE = re.compile(
    r"(?<![\w.])[-+]?(?:(?:\d{1,3}(?:[.,，]\d{3})+)(?:[.,，]\d+)?|\d+(?:[.,，]\d+)?)"
    r"(?:\s*(?:亿元|亿|万元|万|元))?(?![\w.,])",
    re.IGNORECASE,
)
_CONTROL_LABEL_RE = re.compile(
    r"招标控制价|最高限价|控制价|control(?:\s*price)?|ceiling(?:\s*price)?", re.IGNORECASE
)
_EXPLICIT_TOTAL_LABEL_RE = re.compile(
    r"投标总价|投标报价|总报价|含税报价|不含税报价|报价金额|报价合计|含税总价|"
    r"不含税总价|合同总价|项目总价|total(?:\s*price)?|bid[_ -]?price", re.IGNORECASE
)
_GENERIC_TOTAL_LABEL_RE = re.compile(r"合计|金额|amount|price", re.IGNORECASE)
_LABEL_RE = re.compile(
    r"(招标控制价|最高限价|控制价|投标总价|投标报价|总报价|含税报价|不含税报价|报价金额|报价合计|含税总价|"
    r"不含税总价|合同总价|项目总价|合计|金额|total(?:\s*price)?|bid[_ -]?price|amount|price)",
    re.IGNORECASE,
)
_TOTAL_LABEL_RE = re.compile(
    r"招标控制价|最高限价|控制价|投标总价|投标报价|总报价|含税报价|不含税报价|报价金额|报价合计|含税总价|"
    r"不含税总价|合同总价|项目总价|合计|金额|total|bid[_ -]?price|amount|price",
    re.IGNORECASE,
)
_NOISE_FILE_RE = re.compile(r"^(?:~\$|\.)(?:.*)")
_FOREIGN_CURRENCY_RE = re.compile(
    r"(?:[$€£]|\b(?:usd|eur|gbp|jpy|hkd|aud|cad|sgd|krw|inr|rub)\b|"
    r"美元|欧元|英镑|日元|港币|港元|澳元|加元|新加坡元|韩元|卢布)",
    re.IGNORECASE,
)
_NEGATIVE_AMOUNT_RE = re.compile(r"(?<![\w.])[-−]\s*\d")
_QUOTE_NON_AMOUNT_CONTEXT_RE = re.compile(
    r"税率|税额|税点|tax[_ -]?rate|quantity|qty|数量|工程量|"
    r"discount|percent|折扣|下浮率|%",
    re.IGNORECASE,
)

_ITEM_NAME_ALIASES = {"项目名称", "清单名称", "项目", "名称", "name", "item", "description", "清单项目"}
_ITEM_UNIT_ALIASES = {"单位", "计量单位", "unit", "uom"}
_ITEM_SPEC_ALIASES = {"规格", "规格型号", "规格特征", "项目特征", "特征", "spec", "specification", "feature"}
_ITEM_AMOUNT_ALIASES = {"合价", "金额", "小计", "总价", "amount", "total", "合计"}
_ITEM_UNIT_PRICE_ALIASES = {"单价", "unitprice", "price"}
_ITEM_QUANTITY_ALIASES = {"数量", "工程量", "quantity", "qty"}

_FIELD_ALIASES = {
    "author": {"author", "作者", "创建者", "creator", "lastmodifiedby", "最后修改者"},
    "machine_id": {"机器码", "制作机器码", "machinecode", "machine_id", "creator_machine"},
    "mac": {"mac", "mac地址", "mac_address", "网卡mac", "网卡mac地址"},
    "ip": {"ip", "ip地址", "ip_address", "上传ip", "下载ip", "网络地址"},
    "disk_serial": {"硬盘序列号", "硬盘号", "diskserial", "disk_serial"},
    "certificate": {"数字证书", "电子证书", "certificate", "证书编号", "ca证书"},
    "contact": {"联系人", "contact", "contactperson"},
    "contact_phone": {"电话", "联系电话", "联系人电话", "phone", "mobile", "手机号"},
    "legal_representative": {"法定代表人", "法人", "legalrepresentative", "legal_representative"},
    "registered_address": {"注册地址", "注册地", "address", "registeredaddress"},
    "project_manager": {"项目经理", "项目负责人", "projectmanager", "project_manager"},
    "company_uscc": {"统一社会信用代码", "社会信用代码", "uscc", "creditcode"},
}
_CANONICAL_FIELD_BY_ALIAS = {
    re.sub(r"[\s_\-]", "", alias).lower(): field
    for field, aliases in _FIELD_ALIASES.items()
    for alias in aliases
}


def _norm(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().lower()


def _norm_text(value: str) -> str:
    return re.sub(r"\s+", "", value or "").lower()


def _norm_label(value: Any) -> str:
    return re.sub(r"[\s_\-]", "", str(value or "")).lower()


def _json_value(value: Any) -> Any:
    """把 openpyxl 的日期/Decimal 等单元格对象变成可序列化审计值。"""
    if value is None or isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    return str(value)


def _safe_rel(path: Path, root: Path) -> str:
    return path.relative_to(root).as_posix()


def _path_key(path: Path) -> str:
    """比较路径时不解析符号链接，避免把目录外目标当成输入文件。"""
    return os.path.normcase(os.path.abspath(os.fspath(path)))


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def _read_limited(path: Path, limit: int | None = None) -> bytes:
    """在读取前检查普通文件大小，避免把超限文件整体读入内存。"""
    limit = MAX_FILE_BYTES if limit is None else limit
    if path.is_symlink():
        raise ValueError("跳过符号链接，不读取其目标文件")
    try:
        info = os.stat(path, follow_symlinks=False)
    except OSError as exc:
        raise ValueError(f"文件状态读取失败：{exc}") from exc
    if not stat.S_ISREG(info.st_mode):
        raise ValueError("不是普通文件")
    if info.st_size > limit:
        raise ValueError(f"文件超过 {limit // 1024 // 1024} MB 限制")
    try:
        with path.open("rb") as handle:
            raw = handle.read(limit + 1)
    except OSError as exc:
        raise ValueError(f"文件读取失败：{exc}") from exc
    if len(raw) > limit:
        raise ValueError(f"文件超过 {limit // 1024 // 1024} MB 限制")
    return raw


def _error_file_public(path: Path, source: str) -> dict:
    """解析失败时尽量保留文件大小和可安全读取的 SHA-256。"""
    public = {"path": source, "extension": path.suffix.lower(), "parse_status": "ERROR"}
    try:
        info = os.stat(path, follow_symlinks=False)
    except OSError:
        return public
    if not stat.S_ISREG(info.st_mode) or path.is_symlink():
        return public
    public["size_bytes"] = int(info.st_size)
    if info.st_size > MAX_FILE_BYTES:
        return public
    try:
        raw = _read_limited(path)
    except ValueError:
        return public
    public["size_bytes"] = len(raw)
    public["sha256"] = hashlib.sha256(raw).hexdigest()
    return public


def _check_xlsx_archive(path: Path) -> None:
    """在交给 openpyxl 前限制 XLSX ZIP 成员数量和声明的解压大小。"""
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
    except (OSError, zipfile.BadZipFile) as exc:
        raise ValueError(f"XLSX 压缩包无法读取：{exc}") from exc
    if len(members) > MAX_XLSX_ZIP_MEMBERS:
        raise ValueError(f"XLSX 压缩包成员超过 {MAX_XLSX_ZIP_MEMBERS} 个限制")
    total_uncompressed = 0
    for member in members:
        size = member.file_size
        if size < 0 or size > MAX_XLSX_ENTRY_UNCOMPRESSED_BYTES:
            raise ValueError(
                f"XLSX 压缩包成员 {member.filename!r} 解压大小超过 "
                f"{MAX_XLSX_ENTRY_UNCOMPRESSED_BYTES // 1024 // 1024} MB 限制"
            )
        total_uncompressed += size
        if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
            raise ValueError(
                f"XLSX 压缩包声明解压大小超过 {MAX_XLSX_UNCOMPRESSED_BYTES // 1024 // 1024} MB 限制"
            )


def _parse_amount(value: Any) -> float | None:
    if isinstance(value, (dict, list, tuple, set)) or isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) and number >= 0 else None
    text = str(value).strip()
    if _FOREIGN_CURRENCY_RE.search(text):
        return None
    text = text.replace("￥", "").replace("¥", "").replace("人民币", "")
    match = _AMOUNT_RE.search(text)
    if not match:
        return None
    token = match.group(0).replace(" ", "")
    multiplier = 1.0
    for suffix, factor in (("亿元", 100_000_000), ("亿", 100_000_000),
                           ("万元", 10_000), ("万", 10_000), ("元", 1)):
        if token.lower().endswith(suffix):
            token = token[:-len(suffix)]
            multiplier = factor
            break
    sign = "-" if token.startswith("-") else ""
    unsigned = token.lstrip("+-")
    separators = [index for index, char in enumerate(unsigned) if char in ".,，"]
    if separators:
        last = separators[-1]
        last_separator = unsigned[last]
        fraction = unsigned[last + 1:]
        # 两种分隔符同时出现时，最后一个分隔符是小数点：1.234,56 / 1,234.56。
        if "." in unsigned and ("," in unsigned or "，" in unsigned):
            integer = re.sub(r"[.,，]", "", unsigned[:last])
            normalized = f"{sign}{integer}.{fraction}"
        elif len(separators) > 1 and len(set(unsigned[index] for index in separators)) == 1 \
                and len(fraction) == 3 \
                and all(len(part) == 3 for part in re.split(r"[.,，]", unsigned)[1:]):
            normalized = sign + re.sub(r"[.,，]", "", unsigned)
        elif len(separators) > 1:
            integer = re.sub(r"[.,，]", "", unsigned[:last])
            normalized = f"{sign}{integer}.{fraction}"
        elif last_separator in ",，" and len(fraction) == 3:
            normalized = sign + unsigned.replace(last_separator, "")
        elif last_separator == "." and len(fraction) == 3 and not unsigned.startswith("0"):
            # 单个点后三位在金额文本中按千分位处理；0.123 保留为小数。
            normalized = sign + unsigned.replace(".", "")
        else:
            normalized = f"{sign}{unsigned[:last].replace(',', '').replace('，', '')}.{fraction}"
    else:
        normalized = sign + unsigned
    try:
        number = float(normalized) * multiplier
    except ValueError:
        return None
    return number if math.isfinite(number) and number >= 0 else None


def _amounts(value: Any) -> list[float]:
    if isinstance(value, (int, float)):
        parsed = _parse_amount(value)
        return [parsed] if parsed is not None else []
    text = str(value or "")
    if _FOREIGN_CURRENCY_RE.search(text):
        return []
    return [x for x in (_parse_amount(m.group(0)) for m in _AMOUNT_RE.finditer(text)) if x is not None]


def _quote_cell_amounts(value: Any) -> list[float]:
    """只从紧邻报价标签的、没有数量/税率语义的单元格取金额。"""
    text = str(value or "").strip()
    if not text or _QUOTE_NON_AMOUNT_CONTEXT_RE.search(text):
        return []
    return _amounts(text)


def _relative_diff(a: float, b: float) -> float:
    return abs(a - b) / max(abs(a), abs(b), 1.0)


def _quote(label: str, value: float, source: str, locator: str, raw: Any, kind: str = "explicit_total") -> dict:
    return {
        "label": str(label or "报价"),
        "value": round(float(value), 6),
        "source": source,
        "locator": locator,
        "raw": str(raw)[:500],
        "kind": kind,
    }


def _quote_kind(label: Any) -> str:
    text = str(label or "")
    if _CONTROL_LABEL_RE.search(text):
        return "control"
    if _EXPLICIT_TOTAL_LABEL_RE.search(text):
        return "untaxed_total" if "不含税" in text else "explicit_total"
    return "generic_total"


def _quote_rank(quote: dict) -> int:
    if quote.get("kind") == "control" or _CONTROL_LABEL_RE.search(str(quote.get("label") or "")):
        return -1
    label = str(quote.get("label") or "")
    if "不含税" in label:
        return 1
    if "含税" in label:
        return 4
    if _EXPLICIT_TOTAL_LABEL_RE.search(label):
        return 3
    return 0


def _comparison_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().lower()


def _line_item(name: Any, amount: Any, source: str, locator: str, quantity: Any = None,
               unit_price: Any = None, raw: Any = None, unit: Any = None,
               specification: Any = None, feature: Any = None) -> dict | None:
    item_name = str(name or "").strip()
    value = _parse_amount(amount)
    if not item_name or value is None or value == 0:
        return None
    unit_text = str(unit or "").strip()
    specification_text = str(specification or "").strip()
    feature_text = str(feature or "").strip()
    spec_text = specification_text or feature_text
    missing = []
    if not unit_text:
        missing.append("unit")
    if not spec_text:
        missing.append("specification_or_feature")
    # 名称仅作展示/缺口提示；comparison_key 不允许在缺少口径字段时生成。
    comparison_key = None if missing else "|".join(
        (_comparison_text(item_name), _comparison_text(unit_text), _comparison_text(spec_text))
    )
    return {
        "name": item_name[:300],
        "key": _norm(item_name),  # 兼容旧 JSON；比较逻辑只使用 comparison_key
        "name_key": _norm(item_name),
        "unit": unit_text[:100] or None,
        "specification": specification_text[:300] or None,
        "feature": feature_text[:500] or None,
        "comparison_key": comparison_key,
        "comparability_status": "COMPARABLE" if comparison_key else "INSUFFICIENT_DATA",
        "comparability_missing": missing,
        "amount": round(value, 6),
        "quantity": _parse_amount(quantity),
        "unit_price": _parse_amount(unit_price),
        "source": source,
        "locator": locator,
        "raw": str(raw if raw is not None else item_name)[:500],
    }


def _field_name(key: Any) -> str | None:
    return _CANONICAL_FIELD_BY_ALIAS.get(_norm_label(key))


def _metadata_from_pairs(pairs: Iterable[tuple[Any, Any]]) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = {}
    for key, value in pairs:
        field = _field_name(key)
        if not field or value in (None, "", []):
            continue
        text = str(value).strip()
        if not text:
            continue
        out.setdefault(field, []).append({"key": str(key), "value": text[:500]})
    return out


def _merge_metadata(target: dict[str, list[dict]], source: dict[str, list[dict]]) -> None:
    for field, values in source.items():
        target.setdefault(field, []).extend(values)


def _text_quotes(text: str, source: str) -> tuple[list[dict], list[dict]]:
    quotes: list[dict] = []
    items: list[dict] = []
    for index, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        label_match = _LABEL_RE.search(line)
        if label_match:
            tail = line[label_match.end():]
            candidates = _amounts(tail)
            if len(candidates) == 1:
                value = candidates[0]
                label = label_match.group(1)
                quotes.append(_quote(label, value, source, f"第{index}行", line, _quote_kind(label)))
        # 仅把带明显分隔符、且首段像清单名称的行当作明细，降低普通数字误报。
        if any(sep in line for sep in ("\t", "|", ",", "，", ";", "；")):
            cells = [c.strip() for c in re.split(r"\t|\||,|，|;|；", line)]
            if len(cells) >= 2 and not _TOTAL_LABEL_RE.search(cells[0]):
                nums = [_parse_amount(c) for c in cells[1:]]
                nums = [n for n in nums if n is not None]
                if len(nums) >= 1 and not _AMOUNT_RE.fullmatch(cells[0].replace(" ", "")):
                    item = _line_item(cells[0], nums[-1], source, f"第{index}行", raw=line)
                    if item:
                        items.append(item)
    return quotes, items


def _quote_parse_warnings(text: str) -> list[dict]:
    """保留负报价/外币报价的可追溯缺口，不把它们静默当作无报价。"""
    warnings = []
    for index, line in enumerate(text.splitlines(), 1):
        label_match = _LABEL_RE.search(line)
        if not label_match:
            continue
        tail = line[label_match.end():]
        if _NEGATIVE_AMOUNT_RE.search(tail):
            warnings.append({
                "locator": f"第{index}行",
                "reason": "报价金额为负值，未纳入报价比较",
                "raw": line[:500],
            })
        elif _FOREIGN_CURRENCY_RE.search(tail) and _AMOUNT_RE.search(tail):
            warnings.append({
                "locator": f"第{index}行",
                "reason": "报价币种疑似为非人民币，未纳入报价比较",
                "raw": line[:500],
            })
    return warnings


def _rows_from_csv(text: str) -> tuple[list[list[str]], str]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|，；")
    except csv.Error:
        dialect = csv.excel
    rows = list(islice(csv.reader(io.StringIO(text, newline=""), dialect), MAX_ROWS))
    return rows, dialect.delimiter


def _alias_matches(label: str, alias: str) -> bool:
    if alias == label:
        return True
    # 避免把 unit_price 当成 unit、把项目特征当成项目名称。
    if alias in {"unit", "item", "项目"} and any(
        marker in label for marker in ("price", "特征", "feature", "spec")
    ):
        return False
    return alias in label


def _header_candidates(labels: list[str], aliases: set[str]) -> list[int]:
    normalized = [_norm_label(alias) for alias in aliases]
    return [index for index, label in enumerate(labels)
            if any(_alias_matches(label, alias) for alias in normalized)]


def _dict_field(data: dict, aliases: set[str]) -> Any:
    normalized = [_norm_label(alias) for alias in aliases]
    for key, value in data.items():
        label = _norm_label(key)
        if any(_alias_matches(label, alias) for alias in normalized):
            return value
    return None


def _rows_quotes(rows: list[list[Any]], source: str, *, sheet: str | None = None) -> tuple[list[dict], list[dict], dict, dict]:
    quotes: list[dict] = []
    items: list[dict] = []
    metadata: dict[str, list[dict]] = {}
    prefix = f"工作表[{sheet}] " if sheet else ""
    for row_index, row in enumerate(rows, 1):
        cells = ["" if cell is None else str(cell).strip() for cell in row]
        if not any(cells):
            continue
        label_positions = [i for i, cell in enumerate(cells) if _LABEL_RE.search(cell)]
        for pos in label_positions:
            label_match = _LABEL_RE.search(cells[pos])
            tail = cells[pos][label_match.end():] if label_match else ""
            candidates = _amounts(tail) if tail else (
                _quote_cell_amounts(cells[pos + 1]) if pos + 1 < len(cells) else []
            )
            if len(candidates) == 1:
                label = label_match.group(1) if label_match else cells[pos]
                quotes.append(_quote(label, candidates[0], source, f"{prefix}第{row_index}行",
                                     " | ".join(cells), _quote_kind(label)))
        _merge_metadata(metadata, _metadata_from_pairs(zip(cells[::2], cells[1::2])))

    # 识别常见的报价清单表头；明细 amount 优先取合价/金额，其次取单价。
    header_index = None
    name_index = amount_index = quantity_index = unit_index = unit_price_index = None
    specification_index = feature_index = None
    for index, row in enumerate(rows[:30]):
        labels = [_norm_label(cell) for cell in row]
        name_candidates = _header_candidates(labels, _ITEM_NAME_ALIASES)
        amount_candidates = _header_candidates(labels, _ITEM_AMOUNT_ALIASES)
        unit_price_candidates = _header_candidates(labels, _ITEM_UNIT_PRICE_ALIASES)
        unit_candidates = _header_candidates(labels, _ITEM_UNIT_ALIASES)
        quantity_candidates = _header_candidates(labels, _ITEM_QUANTITY_ALIASES)
        specification_candidates = _header_candidates(labels, _ITEM_SPEC_ALIASES)
        feature_candidates = [i for i, label in enumerate(labels) if "特征" in label or "feature" in label]
        if name_candidates and (amount_candidates or unit_price_candidates):
            header_index = index
            name_index = name_candidates[0]
            amount_index = (amount_candidates or unit_price_candidates)[0]
            unit_index = unit_candidates[0] if unit_candidates else None
            unit_price_index = unit_price_candidates[0] if unit_price_candidates else None
            quantity_index = quantity_candidates[0] if quantity_candidates else None
            specification_index = specification_candidates[0] if specification_candidates else None
            feature_index = feature_candidates[0] if feature_candidates else None
            break
    if header_index is not None and name_index is not None and amount_index is not None:
        for row_index, row in enumerate(rows[header_index + 1:], header_index + 2):
            cells = ["" if cell is None else str(cell).strip() for cell in row]
            if len(cells) <= max(name_index, amount_index) or _TOTAL_LABEL_RE.search(" ".join(cells[:2])):
                continue
            item = _line_item(
                cells[name_index], cells[amount_index], source, f"{prefix}第{row_index}行",
                cells[quantity_index] if quantity_index is not None and quantity_index < len(cells) else None,
                cells[unit_price_index] if unit_price_index is not None and unit_price_index < len(cells) else None,
                " | ".join(cells),
                cells[unit_index] if unit_index is not None and unit_index < len(cells) else None,
                cells[specification_index] if specification_index is not None and specification_index < len(cells) else None,
                cells[feature_index] if feature_index is not None and feature_index < len(cells) else None,
            )
            if item:
                items.append(item)
    structure = {
        "row_count": len(rows),
        "column_count": max((len(row) for row in rows), default=0),
        "header": [_json_value(value) for value in (
            rows[header_index] if header_index is not None else (rows[0] if rows else [])
        )],
    }
    return quotes, items, structure, metadata


def _json_walk(value: Any, source: str, path: str = "$", quotes: list | None = None,
               items: list | None = None, metadata: dict | None = None, *,
               depth: int = 0, budget: dict[str, int] | None = None) -> tuple[list, list, dict]:
    if depth > MAX_JSON_DEPTH:
        raise ValueError(f"JSON 嵌套深度超过 {MAX_JSON_DEPTH} 限制")
    budget = budget if budget is not None else {"nodes": 0}
    budget["nodes"] += 1
    if budget["nodes"] > MAX_JSON_NODES:
        raise ValueError(f"JSON 节点数超过 {MAX_JSON_NODES} 限制")
    quotes = quotes if quotes is not None else []
    items = items if items is not None else []
    metadata = metadata if metadata is not None else {}
    if isinstance(value, dict):
        _merge_metadata(metadata, _metadata_from_pairs(value.items()))
        name = _dict_field(value, _ITEM_NAME_ALIASES)
        amount = _dict_field(value, _ITEM_AMOUNT_ALIASES)
        if amount is None:
            amount = _dict_field(value, _ITEM_UNIT_PRICE_ALIASES)
        parsed_amount = _parse_amount(amount)
        if name is not None and parsed_amount is not None:
            unit = _dict_field(value, _ITEM_UNIT_ALIASES)
            specification = _dict_field(value, _ITEM_SPEC_ALIASES)
            feature = _dict_field(value, {"特征", "项目特征", "feature"})
            item = _line_item(name, parsed_amount, source, f"JSON {path}", raw=value,
                              unit=unit, specification=specification, feature=feature)
            if item:
                items.append(item)
        for key, child in value.items():
            # total/bid_price 可能是包含 value、currency 等字段的复合对象；
            # 复合对象本身不是报价，继续递归但不把其字符串表示当数字解析。
            if (_CONTROL_LABEL_RE.search(str(key)) or _EXPLICIT_TOTAL_LABEL_RE.search(str(key))) \
                    and not isinstance(child, (dict, list, tuple, set)):
                amount = _parse_amount(child)
                if amount is not None:
                    quotes.append(_quote(str(key), amount, source, f"JSON {path}.{key}", child,
                                         _quote_kind(key)))
            _json_walk(child, source, f"{path}.{key}", quotes, items, metadata,
                       depth=depth + 1, budget=budget)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            _json_walk(child, source, f"{path}[{index}]", quotes, items, metadata,
                       depth=depth + 1, budget=budget)
    return quotes, items, metadata


def _dedupe_quotes(quotes: list[dict]) -> list[dict]:
    seen: set[tuple] = set()
    out = []
    for quote in quotes:
        key = (quote["source"], quote["locator"], quote["value"], quote["label"])
        if key not in seen:
            seen.add(key)
            out.append(quote)
    return out


def _dedupe_items(items: list[dict]) -> list[dict]:
    seen: set[tuple] = set()
    out = []
    for item in items:
        key = (
            item["source"], item["locator"], item["name_key"], item.get("unit"),
            item.get("specification"), item.get("feature"), item["amount"],
        )
        if key not in seen:
            seen.add(key)
            out.append(item)
    return out


def _primary_quote(quotes: list[dict]) -> dict | None:
    candidates = [quote for quote in quotes if _quote_rank(quote) >= 0]
    if not candidates:
        return None
    best_rank = max(_quote_rank(quote) for quote in candidates)
    best = [quote for quote in candidates if _quote_rank(quote) == best_rank]
    values = {float(quote["value"]) for quote in best}
    # 同优先级出现不同金额时不猜主报价；即使同值，也保留第一份可追溯来源。
    return best[0] if len(values) == 1 else None


def _file_metadata(path: Path, source: str) -> tuple[dict, str, list[dict], list[dict], dict, dict[str, list[dict]], str | None]:
    """读取一个文件，返回 public meta、compare text、quotes、items、structure、metadata、error。"""
    raw = _read_limited(path)
    digest = hashlib.sha256(raw).hexdigest()
    suffix = path.suffix.lower()
    quotes: list[dict] = []
    items: list[dict] = []
    structure: dict = {}
    metadata: dict[str, list[dict]] = {}
    parse_warnings: list[dict] = []
    if suffix in {".txt", ".md"}:
        text = _decode(raw)
        quotes, items = _text_quotes(text, source)
        structure = {"line_count": len(text.splitlines()), "column_count": 1}
    elif suffix == ".csv":
        text = _decode(raw)
        rows, delimiter = _rows_from_csv(text)
        quotes, items, structure, metadata = _rows_quotes(rows, source)
        structure["delimiter"] = delimiter
    elif suffix == ".json":
        text = _decode(raw)
        data = json.loads(text)
        quotes, items, metadata = _json_walk(data, source)
        structure = {"json_type": type(data).__name__, "top_level_keys": list(data)[:50] if isinstance(data, dict) else []}
        text = json.dumps(data, ensure_ascii=False, sort_keys=True)
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - package declares dependency
            raise ValueError("读取 xlsx 需要 openpyxl") from exc
        _check_xlsx_archive(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet_summaries = []
        text_parts = []
        read_cell_count = 0
        read_truncated = False
        try:
            for sheet in workbook.worksheets:
                rows = []
                declared_rows = max(1, int(sheet.max_row or 1))
                declared_columns = max(1, int(sheet.max_column or 1))
                remaining_cells = MAX_XLSX_AUDIT_CELLS - read_cell_count
                if remaining_cells > 0:
                    read_columns = min(declared_columns, remaining_cells)
                    read_rows = min(
                        declared_rows, MAX_ROWS,
                        max(1, remaining_cells // read_columns),
                    )
                    sheet_truncated = (
                        read_rows < declared_rows or read_columns < declared_columns
                        or (read_rows == MAX_ROWS and declared_rows > MAX_ROWS)
                    )
                    for row in sheet.iter_rows(
                            max_row=read_rows, max_col=read_columns, values_only=True):
                        values = list(row)
                        rows.append(values)
                        read_cell_count += len(values)
                    read_truncated = read_truncated or sheet_truncated
                else:
                    sheet_truncated = True
                    read_truncated = True
                q, i, s, row_metadata = _rows_quotes(rows, source, sheet=sheet.title)
                quotes.extend(q)
                items.extend(i)
                _merge_metadata(metadata, row_metadata)
                sheet_summaries.append({"name": sheet.title, "read_truncated": sheet_truncated, **s})
                text_parts.append(f"[SHEET:{sheet.title}]\n" + "\n".join(" | ".join("" if x is None else str(x) for x in row) for row in rows))
            # read_only 模式适合取显示值，但不会保留隐藏行、筛选和公式类型；
            # 只有显示值读取未触及审计上限时，才做一次普通读取补充元数据；
            # 否则不再为稀疏超宽/超长工作表加载完整对象。
            audit_cell_count = 0
            audit_truncated = read_truncated
            if not read_truncated:
                audit_workbook = load_workbook(path, read_only=False, data_only=False)
                try:
                    for summary in sheet_summaries:
                        audit_sheet = audit_workbook[summary["name"]]
                        summary["hidden_row_count"] = sum(
                            1 for dimension in audit_sheet.row_dimensions.values() if dimension.hidden
                        )
                        summary["filtered_range"] = audit_sheet.auto_filter.ref
                        summary["sheet_state"] = audit_sheet.sheet_state
                        summary["merged_range_count"] = len(audit_sheet.merged_cells.ranges)
                        formula_count = 0
                        for row in audit_sheet.iter_rows():
                            for cell in row:
                                if audit_cell_count >= MAX_XLSX_AUDIT_CELLS:
                                    audit_truncated = True
                                    break
                                audit_cell_count += 1
                                if cell.data_type == "f":
                                    formula_count += 1
                            if audit_truncated:
                                break
                        summary["formula_count"] = formula_count
                        summary["audit_truncated"] = audit_truncated
                finally:
                    audit_workbook.close()
            else:
                for summary in sheet_summaries:
                    summary["hidden_row_count"] = None
                    summary["filtered_range"] = None
                    summary["sheet_state"] = None
                    summary["merged_range_count"] = None
                    summary["formula_count"] = None
                    summary["audit_truncated"] = True
            formula_count = sum(
                int(summary.get("formula_count") or 0) for summary in sheet_summaries
            )
            if formula_count:
                parse_warnings.append({
                    "locator": "XLSX 工作表",
                    "reason": f"XLSX 含 {formula_count} 个公式；程序不计算公式，报价可能依赖缓存显示值，需在 Excel/WPS 重算后人工复核",
                    "raw": path.name,
                })
            properties = workbook.properties
            _merge_metadata(metadata, _metadata_from_pairs((("author", properties.creator), ("lastModifiedBy", properties.lastModifiedBy))))
        finally:
            workbook.close()
        structure = {
            "sheet_count": len(sheet_summaries), "sheets": sheet_summaries,
            "audit_cell_count": audit_cell_count, "audit_truncated": audit_truncated,
        }
        text = "\n".join(text_parts)
    else:
        return {"path": source, "extension": suffix, "size_bytes": len(raw), "sha256": digest,
                "parse_status": "UNSUPPORTED"}, "", [], [], {}, {}, f"不支持的文件类型：{suffix or '无扩展名'}"
    compare_text = text[:MAX_TEXT_FOR_COMPARE]
    public = {
        "path": source,
        "extension": suffix,
        "size_bytes": len(raw),
        "sha256": digest,
        "normalized_sha256": hashlib.sha256(_norm_text(text).encode("utf-8")).hexdigest(),
        "parse_status": "OK",
        "text_chars": len(text),
        "quote_count": len(_dedupe_quotes(quotes)),
        "line_item_count": len(_dedupe_items(items)),
        "structure": structure,
        "metadata_fields": sorted(metadata),
    }
    parse_warnings.extend(_quote_parse_warnings(text))
    if parse_warnings:
        public["parse_warnings"] = parse_warnings
    return public, compare_text, _dedupe_quotes(quotes), _dedupe_items(items), structure, metadata, None


def _new_signal(code: str, title: str, scope: str, description: str, evidence: list[dict], *, level: str = "中",
                action: str = "人工复核相关原始文件、电子投标平台日志及业务口径") -> dict:
    return {
        "code": code,
        "rule_id": RULE_IDS.get(code),
        "kind": code.lower(),
        "level": level,
        "title": title,
        "scope": scope,
        "description": description,
        "evidence": evidence,
        "manual_action": action,
        "auto_conclusion": False,
    }


def _pairwise(values: list[str]) -> Iterable[tuple[str, str]]:
    for index, left in enumerate(values):
        for right in values[index + 1:]:
            yield left, right


def _bidder_pair_maps(bidders: list[dict]) -> dict[str, dict]:
    return {bidder["name"]: bidder for bidder in bidders}


def _compare_quotes(bidders: list[dict], signals: list[dict]) -> None:
    quotes = {bidder["name"]: bidder.get("primary_quote") for bidder in bidders if bidder.get("primary_quote")}
    names = list(quotes)
    for left, right in _pairwise(names):
        a, b = quotes[left], quotes[right]
        diff = _relative_diff(a["value"], b["value"])
        if diff <= THRESHOLDS["near_quote_relative_diff"]:
            signals.append(_new_signal(
                "QUOTE_NEAR_MATCH", "投标总报价高度接近", f"{left} ↔ {right}",
                f"两份投标资料主报价分别为 {a['value']:,.2f} 与 {b['value']:,.2f}，相对差异约 {diff:.2%}。该数值模式可能有正常报价口径解释，不能单独认定违法。",
                [a, b], level="高"))
    if len(quotes) >= 3:
        values = list(quotes.values())
        numbers = [float(x["value"]) for x in values]
        average = mean(numbers)
        cv = pstdev(numbers) / average if average else 0.0
        if cv <= THRESHOLDS["low_dispersion_cv"]:
            signals.append(_new_signal(
                "QUOTE_LOW_DISPERSION", "多家投标报价离散度偏低", "项目全部已识别报价",
                f"{len(numbers)} 家主报价变异系数约 {cv:.2%}，低于预警阈值 {THRESHOLDS['low_dispersion_cv']:.2%}；需先核对控制价、清单口径、四舍五入和报价策略。",
                list(quotes.values()), level="中"))
        middle = median(numbers)
        for name, quote in quotes.items():
            deviation = _relative_diff(quote["value"], middle)
            if deviation >= THRESHOLDS["outlier_relative_diff"]:
                signals.append(_new_signal(
                    "QUOTE_OUTLIER", "单家投标报价相对中位数离群", name,
                    f"主报价 {quote['value']:,.2f} 与项目报价中位数 {middle:,.2f} 的相对差异约 {deviation:.2%}。离群本身可能由方案、范围或计价口径造成，需人工核对。",
                    [quote], level="中"))


def _compare_line_items(bidders: list[dict], signals: list[dict]) -> None:
    maps: dict[str, dict[str, list[dict]]] = {}
    for bidder in bidders:
        grouped: dict[str, list[dict]] = {}
        for item in bidder.get("line_items", []):
            # 仅使用带单位和规格/特征的显式口径键；缺字段的项目保留在结果中，
            # 但不得靠名称自动对齐或触发同步报价预警。
            key = item.get("comparison_key")
            if key and item.get("comparability_status") == "COMPARABLE":
                grouped.setdefault(key, []).append(item)
        maps[bidder["name"]] = grouped
    names = list(maps)
    for left, right in _pairwise(names):
        # 同一投标人内同键重复也属于歧义，不自动合并。
        common = sorted(
            key for key in set(maps[left]) & set(maps[right])
            if len(maps[left][key]) == 1 and len(maps[right][key]) == 1
        )
        if len(common) < THRESHOLDS["min_common_line_items"]:
            continue
        matched = []
        ratios = []
        for key in common:
            a, b = maps[left][key][0], maps[right][key][0]
            diff = _relative_diff(a["amount"], b["amount"])
            if diff <= THRESHOLDS["line_item_match_relative_diff"]:
                matched.append((a, b))
            if b["amount"]:
                ratios.append(a["amount"] / b["amount"])
        ratio_cv = pstdev(ratios) / mean(ratios) if len(ratios) >= 3 and mean(ratios) else None
        enough_matches = len(matched) >= max(THRESHOLDS["min_common_line_items"], math.ceil(len(common) * 0.6))
        same_scale = ratio_cv is not None and ratio_cv <= THRESHOLDS["low_dispersion_cv"]
        if enough_matches or same_scale:
            evidence = []
            for key in common[:20]:
                evidence.extend([maps[left][key][0], maps[right][key][0]])
            signals.append(_new_signal(
                "SYNCHRONIZED_LINE_ITEMS", "多个清单项报价呈同步或同尺度变化", f"{left} ↔ {right}",
                f"两家共有 {len(common)} 个可比清单项，其中 {len(matched)} 个金额高度接近"
                + (f"，共有项报价比的变异系数约 {ratio_cv:.2%}" if ratio_cv is not None else "")
                + "；需排除统一计价规则、清单版本和单位换算造成的共同原因。",
                evidence, level="高"))


def _compare_text_and_structure(bidders: list[dict], signals: list[dict]) -> None:
    # ponytail: pairwise comparison is O(n²) and SequenceMatcher is O(n²) worst-case;
    # keep the bounded offline scan simple, upgrade to blocking/MinHash if scale requires it.
    for left, right in _pairwise([b["name"] for b in bidders]):
        a = next(b for b in bidders if b["name"] == left)
        b = next(b for b in bidders if b["name"] == right)
        best_similarity = (0.0, None, None)
        structure_a = set()
        structure_b = set()
        for file in a["files"]:
            structure_a.add((file["extension"], Path(file["path"]).name.lower()))
        for file in b["files"]:
            structure_b.add((file["extension"], Path(file["path"]).name.lower()))
        shared = structure_a & structure_b
        union = structure_a | structure_b
        structure_score = len(shared) / len(union) if union else 0.0
        if structure_score >= THRESHOLDS["structure_similarity"] and len(shared) >= 2:
            signals.append(_new_signal(
                "STRUCTURE_SIMILARITY", "投标文件结构/文件名组合高度相似", f"{left} ↔ {right}",
                f"两家文件结构组合交并比约 {structure_score:.2%}，共有 {len(shared)} 个文件名/类型组合；相同模板可能有合理来源，需结合内容和形成过程复核。",
                [{"bidder": left, "files": a["files"]}, {"bidder": right, "files": b["files"]}], level="中"))
        for file_a in a["_internal_files"]:
            for file_b in b["_internal_files"]:
                if (file_a["public"]["sha256"] == file_b["public"]["sha256"]
                        and min(len(file_a["text"]), len(file_b["text"])) >= THRESHOLDS["min_text_chars"]):
                    signals.append(_new_signal(
                        "TEXT_EXACT_MATCH", "投标文件字节内容完全一致", f"{left} ↔ {right}",
                        "两份文件 SHA-256 一致，表明字节内容相同；仍需核对模板、招标文件统一附件及合法共享文件来源。",
                        [{"bidder": left, **file_a["public"]}, {"bidder": right, **file_b["public"]}], level="高"))
                elif (file_a["public"].get("normalized_sha256") == file_b["public"].get("normalized_sha256")
                      and min(len(file_a["text"]), len(file_b["text"])) >= THRESHOLDS["min_text_chars"]):
                    signals.append(_new_signal(
                        "TEXT_EXACT_MATCH", "投标文件去空白文本内容一致", f"{left} ↔ {right}",
                        "两份文件去除空白后的文本摘要一致，需人工排除统一模板、公告附件或同一原始文件合法复用。",
                        [{"bidder": left, **file_a["public"]}, {"bidder": right, **file_b["public"]}], level="高"))
                elif file_a["public"]["extension"] == file_b["public"]["extension"] and min(len(file_a["text"]), len(file_b["text"])) >= THRESHOLDS["min_text_chars"]:
                    score = difflib.SequenceMatcher(None, file_a["text"][:MAX_TEXT_FOR_COMPARE], file_b["text"][:MAX_TEXT_FOR_COMPARE]).ratio()
                    if score > best_similarity[0]:
                        best_similarity = (score, file_a, file_b)
        if best_similarity[1] is not None and best_similarity[0] >= THRESHOLDS["text_similarity"]:
            score, file_a, file_b = best_similarity
            signals.append(_new_signal(
                "TEXT_HIGH_SIMILARITY", "投标文件文本相似度较高", f"{left} ↔ {right}",
                f"同类型文件文本 SequenceMatcher 相似度约 {score:.2%}，需核对招标文件统一格式、固定条款、复制粘贴来源和形成时间。",
                [{"bidder": left, **file_a["public"]}, {"bidder": right, **file_b["public"]}], level="中"))


def _compare_metadata(bidders: list[dict], signals: list[dict]) -> None:
    maps: dict[str, dict[str, set[str]]] = {}
    for bidder in bidders:
        by_field: dict[str, set[str]] = {}
        for field, values in bidder.get("metadata", {}).items():
            if field not in _ELECTRONIC_METADATA_FIELDS:
                continue
            by_field.setdefault(field, set()).update(_norm(v["value"]) for v in values if v.get("value"))
        maps[bidder["name"]] = by_field
    for left, right in _pairwise(list(maps)):
        for field in sorted(set(maps[left]) & set(maps[right])):
            shared = maps[left][field] & maps[right][field]
            if not shared:
                continue
            signals.append(_new_signal(
                "METADATA_MATCH", "投标资料元数据/主体字段存在相同值", f"{left} ↔ {right}",
                f"字段 {field} 出现相同值（值已在证据中保留）；单个作者、IP、MAC、联系人或地址相同均不能单独作出违法判断。",
                [{"bidder": left, "field": field, "values": sorted(shared)},
                 {"bidder": right, "field": field, "values": sorted(shared)}], level="中"))


def _relation_rows(path: Path, raw: bytes | None = None) -> tuple[list[dict], str | None]:
    try:
        raw = raw if raw is not None else _read_limited(path)
    except ValueError as exc:
        return [], str(exc)
    text = _decode(raw)
    if path.suffix.lower() == ".json":
        try:
            data = json.loads(text)
        except json.JSONDecodeError as exc:
            return [], f"JSON 解析失败：{exc.msg}"
        try:
            _json_walk(data, str(path))
        except ValueError as exc:
            return [], f"JSON 解析受限：{exc}"
        if isinstance(data, dict):
            data = data.get("relations", data.get("clues", data.get("data", [data])))
        if not isinstance(data, list):
            return [], "关联线索 JSON 须为数组或包含 relations/clues/data 数组"
        return [dict(row) for row in data if isinstance(row, dict)], None
    try:
        rows, _ = _rows_from_csv(text)
        if not rows:
            return [], None
        headers = rows[0]
        return [dict(zip(headers, row)) for row in rows[1:]], None
    except (csv.Error, ValueError) as exc:
        return [], f"CSV 解析失败：{exc}"


def _pick(row: dict, aliases: set[str]) -> Any:
    for key, value in row.items():
        if _norm_label(key) in {_norm_label(alias) for alias in aliases}:
            return value
    return None


def _add_relation_signals(path: Path | None, bidder_names: set[str], signals: list[dict], result: dict) -> None:
    if not path:
        return
    if path.is_symlink():
        result["skipped_files"].append({"path": str(path), "reason": "SYMLINK_SKIPPED"})
        return
    if not path.is_file():
        result["parse_errors"].append({"path": str(path), "error": "关联线索文件不存在"})
        return
    try:
        raw = _read_limited(path)
    except ValueError as exc:
        result["parse_errors"].append({"path": str(path), "error": str(exc)})
        return
    result["relation_source"] = {
        "path": str(path), "size_bytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()
    }
    observed_bytes = result["scan"]["total_bytes"] + len(raw)
    if observed_bytes > MAX_INPUT_BYTES:
        result["scan"]["truncated"] = True
        result["scan"]["budget_errors"].append({
            "kind": "bytes", "limit": MAX_INPUT_BYTES, "observed": observed_bytes,
        })
        result["scan_errors"].append({
            "path": str(path), "error": f"输入资料总字节数超过限制：{MAX_INPUT_BYTES}"
        })
        return
    result["scan"]["total_bytes"] = observed_bytes
    rows, error = _relation_rows(path, raw)
    if error:
        result["parse_errors"].append({"path": str(path), "error": error})
        return
    result["relation_clues"] = []
    aliases_a = {"bidder_a", "company_a", "party_a", "投标人a", "企业a", "甲方", "name_a"}
    aliases_b = {"bidder_b", "company_b", "party_b", "投标人b", "企业b", "乙方", "name_b"}
    relation_aliases = {"relation", "relationship", "关联类型", "关联关系", "线索", "说明", "type"}
    source_aliases = {"source", "来源", "evidence", "证据", "依据", "source_note"}
    for row_number, row in enumerate(rows, 2):
        left = str(_pick(row, aliases_a) or "").strip()
        right = str(_pick(row, aliases_b) or "").strip()
        relation = str(_pick(row, relation_aliases) or "用户提供的关联关系线索").strip()
        source = str(_pick(row, source_aliases) or "未注明").strip()
        clue = {"row": row_number, "bidder_a": left, "bidder_b": right, "relation": relation,
                "source": source, "raw": row}
        result["relation_clues"].append(clue)
        left_match = left in bidder_names
        right_match = right in bidder_names
        if left_match or right_match:
            scope = f"{left} ↔ {right}" if left and right else (left or right)
            signals.append(_new_signal(
                "LOCAL_RELATION_CLUE", "本地资料提供投标人关联关系线索", scope,
                f"依据用户提供的本地线索：{relation}；来源/依据：{source}。该线索仅进入人工复核清单，不自动认定关联关系、串通投标或投标无效。",
                [{"source_file": path.name, "row": row_number, "clue": clue}], level="高"))


def _scan_input_files(root: Path, result: dict, excluded_paths: set[Path] | None = None) -> list[Path]:
    excluded_paths = {_path_key(path) for path in (excluded_paths or set())}
    scan = result["scan"]
    paths: list[Path] = []

    def record_skip(path: Path, reason: str) -> None:
        result["skipped_files"].append({"path": _safe_rel(path, root), "reason": reason})
        if reason == "SYMLINK_SKIPPED":
            scan["symlink_count"] += 1

    def stop_for_budget(path: Path, kind: str, limit: int, observed: int) -> None:
        scan["truncated"] = True
        detail = {"kind": kind, "limit": limit, "observed": observed}
        scan["budget_errors"].append(detail)
        label = "文件数" if kind == "file_count" else "总字节数"
        result["scan_errors"].append({
            "path": _safe_rel(path, root),
            "error": f"输入目录超过{label}限制：{limit}",
        })

    def record_walk_error(exc: OSError) -> None:
        raw_path = getattr(exc, "filename", None)
        path = Path(raw_path) if raw_path else root
        try:
            display_path = _safe_rel(path, root)
        except (TypeError, ValueError):
            display_path = str(path)
        result["scan_errors"].append({
            "path": display_path,
            "error": f"目录扫描失败：{exc}",
        })

    for directory, dirnames, filenames in os.walk(
            root, topdown=True, followlinks=False, onerror=record_walk_error):
        directory_path = Path(directory)
        for dirname in sorted(dirnames):
            candidate = directory_path / dirname
            if os.path.islink(candidate):
                record_skip(candidate, "SYMLINK_SKIPPED")
        dirnames[:] = sorted(
            dirname for dirname in dirnames
            if not dirname.startswith(".") and not os.path.islink(directory_path / dirname)
        )
        for filename in sorted(filenames):
            path = directory_path / filename
            if os.path.islink(path):
                record_skip(path, "SYMLINK_SKIPPED")
                continue
            if _path_key(path) in excluded_paths or any(part.startswith(".") for part in path.relative_to(root).parts):
                continue
            if _NOISE_FILE_RE.match(path.name):
                continue
            try:
                file_stat = os.stat(path, follow_symlinks=False)
            except OSError as exc:
                result["scan_errors"].append({
                    "path": _safe_rel(path, root), "error": f"文件状态读取失败：{exc}"
                })
                continue
            if not stat.S_ISREG(file_stat.st_mode):
                continue
            if scan["regular_file_count"] >= MAX_INPUT_FILES:
                stop_for_budget(path, "file_count", MAX_INPUT_FILES, scan["regular_file_count"] + 1)
                return paths
            size = file_stat.st_size
            if scan["total_bytes"] + size > MAX_INPUT_BYTES:
                stop_for_budget(path, "bytes", MAX_INPUT_BYTES, scan["total_bytes"] + size)
                return paths
            scan["regular_file_count"] += 1
            scan["total_bytes"] += size
            paths.append(path)
    return paths


def _load_bidder_files(root: Path, result: dict, excluded_paths: set[Path] | None = None) -> list[dict]:
    files_by_bidder: dict[str, list[Path]] = {}
    for path in _scan_input_files(root, result, excluded_paths):
        relative_parts = path.relative_to(root).parts
        if len(relative_parts) >= 2:
            bidder = relative_parts[0].strip() or "未命名投标人"
        else:
            stem = path.stem.split("__", 1)[0].strip()
            bidder = stem or root.name or "未命名投标人"
        files_by_bidder.setdefault(bidder, []).append(path)
    bidders = []
    for name in sorted(files_by_bidder):
        public_files = []
        internal_files = []
        quotes = []
        items = []
        metadata: dict[str, list[dict]] = {}
        for path in files_by_bidder[name]:
            relative = _safe_rel(path, root)
            try:
                public, text, q, i, structure, file_meta, error = _file_metadata(path, relative)
            except Exception as exc:
                public = _error_file_public(path, relative)
                text, q, i, structure, file_meta, error = "", [], [], {}, {}, str(exc)
            public_files.append(public)
            for warning in public.get("parse_warnings", []):
                result["parse_warnings"].append({
                    "path": relative,
                    "bidder": name,
                    "locator": warning["locator"],
                    "error": warning["reason"],
                    "raw": warning["raw"],
                })
            if error and public.get("parse_status") != "UNSUPPORTED":
                parse_error = {"path": relative, "bidder": name, "error": error}
                for key in ("size_bytes", "sha256"):
                    if key in public:
                        parse_error[key] = public[key]
                result["parse_errors"].append(parse_error)
            if public.get("parse_status") == "UNSUPPORTED":
                result["unsupported_files"].append(public)
            if public.get("parse_status") == "OK":
                internal_files.append({"public": public, "text": text, "structure": structure})
                quotes.extend(q)
                items.extend(i)
                _merge_metadata(metadata, file_meta)
        quotes = _dedupe_quotes(quotes)
        items = _dedupe_items(items)
        bidders.append({
            "name": name,
            "files": public_files,
            "quotes": quotes,
            "primary_quote": _primary_quote(quotes),
            "line_items": items,
            "line_item_count": len(items),
            "comparable_line_item_count": sum(
                1 for item in items if item.get("comparability_status") == "COMPARABLE"
            ),
            "insufficient_line_item_count": sum(
                1 for item in items if item.get("comparability_status") != "COMPARABLE"
            ),
            "metadata": metadata,
            "_internal_files": internal_files,
        })
    return bidders


def review_directory(input_dir: str | Path, *, project: str = "", relations: str | Path | None = None) -> dict:
    """审查一个本地目录，返回 JSON 可序列化结果。

    目录约定：一级子目录名为投标人；根目录文件可用 bidder__filename.ext，
    没有前缀时归入根目录名对应的单一投标人。只扫描支持的五类文件及其解析错误。
    """
    root = Path(input_dir).expanduser()
    if root.is_symlink():
        raise ValueError(f"输入目录不能是符号链接：{root}")
    if not root.exists() or not root.is_dir():
        raise ValueError(f"输入目录不存在或不是目录：{root}")
    root = root.resolve()
    result = {
        "product": "投标审查器",
        "review_type": "offline_multi_bidder_file_review",
        "project": project.strip(),
        "input_dir": str(root),
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "config": {
            "supported_extensions": sorted(SUPPORTED_EXTENSIONS),
            "thresholds": THRESHOLDS,
            "limits": {
                "max_file_bytes": MAX_FILE_BYTES,
                "max_input_files": MAX_INPUT_FILES,
                "max_input_bytes": MAX_INPUT_BYTES,
                "max_xlsx_zip_members": MAX_XLSX_ZIP_MEMBERS,
                "max_xlsx_uncompressed_bytes": MAX_XLSX_UNCOMPRESSED_BYTES,
                "max_xlsx_entry_uncompressed_bytes": MAX_XLSX_ENTRY_UNCOMPRESSED_BYTES,
                "max_json_depth": MAX_JSON_DEPTH,
                "max_json_nodes": MAX_JSON_NODES,
                "max_xlsx_audit_cells": MAX_XLSX_AUDIT_CELLS,
            },
            "enabled_rule_ids": sorted(set(RULE_IDS.values())),
            "rule_id_scheme": "F=报价模式，E=电子/元数据，S=文件相似，P=主体/关系线索；未列入 enabled_rule_ids 的编号当前未启用",
        },
        "legal_notice": {
            "conclusion": "仅输出风险预警和人工复核线索，不构成串通投标、违法、资格不合格或投标无效认定。",
            "scope": "本地依法取得并有权使用的投标文件、电子投标平台导出信息和关联关系线索；检测结果受资料完整性、版本、单位和口径影响。",
            "sources": [
                {"name": "中华人民共和国招标投标法（市场监管总局公布文本）", "url": "https://www.samr.gov.cn/zw/zfxxgk/fdzdgknr/bgt/art/2023/art_1f79dd79321441a0831f3aed697b4535.html"},
                {"name": "中华人民共和国招标投标法实施条例（司法部行政法规库）", "url": "https://xzfg.moj.gov.cn/front/law/detail?LawID=1154"},
                {"name": "电子招标投标办法（国家发展改革委令第20号）", "url": "https://www.ndrc.gov.cn/xxgk/zcfb/fzggwl/201302/t20130220_960752.html?state=123"},
                {"name": "发改法规规〔2022〕1117号意见", "url": "https://www.ndrc.gov.cn/xxgk/zcfb/ghxwj/202208/t20220801_1332495.html"},
                {"name": "浙江省人民政府关于进一步构建规范有序招标投标市场的若干意见（浙政发〔2024〕17号）", "url": "https://zjjcmspublic.oss-cn-hangzhou-zwynet-d01-a.internet.cloud.zj.gov.cn/jcms_files/jcms1/web3241/site/attach/0/75fc778458644eafab722be1603649fc.pdf"},
                {"name": "杭建市〔2020〕190号电子招标投标管理办法", "url": "https://zfgb.hangzhou.gov.cn/upload/default/bigfile/2025/06/09/20250609_58d18709078bc6deaff0044a11891e10.pdf"},
                {"name": "财政部令第87号（仅政府采购适用）", "url": "https://tfs.mof.gov.cn/caizhengbuling/201707/t20170718_2652603.htm", "scope": "仅政府采购"},
                {"name": "浙财采监〔2025〕2号（仅政府采购适用）", "url": "https://czj.hangzhou.gov.cn/art/2025/5/13/art_1655737_59023282.html", "scope": "仅政府采购"},
            ],
        },
        "bidders": [],
        "signals": [],
        "relation_clues": [],
        "unsupported_files": [],
        "parse_errors": [],
        "parse_warnings": [],
        "skipped_files": [],
        "scan_errors": [],
        "scan": {
            "regular_file_count": 0,
            "total_bytes": 0,
            "symlink_count": 0,
            "truncated": False,
            "budget_errors": [],
        },
    }
    relation_path = Path(relations).expanduser().absolute() if relations else None
    bidders = _load_bidder_files(root, result, {relation_path} if relation_path else set())
    # 内部文本只用于本轮比较，不进入最终 JSON，避免报告意外携带整份投标文件。
    result["bidders"] = [{k: v for k, v in bidder.items() if k != "_internal_files"} for bidder in bidders]
    signals: list[dict] = result["signals"]
    _compare_quotes(bidders, signals)
    _compare_line_items(bidders, signals)
    _compare_text_and_structure(bidders, signals)
    _compare_metadata(bidders, signals)
    _add_relation_signals(relation_path, {b["name"] for b in bidders}, signals, result)
    apply_pattern_signals(result, bidders, signals)
    for index, signal in enumerate(signals, 1):
        signal["id"] = f"S{index:03d}"
    all_items = [item for bidder in bidders for item in bidder.get("line_items", [])]
    comparable_items = [item for item in all_items if item.get("comparability_status") == "COMPARABLE"]
    result["summary"] = {
        "bidder_count": len(bidders),
        "file_count": sum(len(b["files"]) for b in bidders),
        "parsed_file_count": sum(sum(1 for f in b["files"] if f.get("parse_status") == "OK") for b in bidders),
        "quote_count": sum(len(b["quotes"]) for b in bidders),
        "signal_count": len(signals),
        "unsupported_file_count": len(result["unsupported_files"]),
        "parse_error_count": len(result["parse_errors"]),
        "parse_warning_count": len(result["parse_warnings"]),
        "skipped_file_count": len(result["skipped_files"]),
        "scan_error_count": len(result["scan_errors"]),
        "line_item_count": len(all_items),
        "comparable_line_item_count": len(comparable_items),
        "insufficient_line_item_count": len(all_items) - len(comparable_items),
        "line_item_comparability": {
            "requires_unit_and_specification_or_feature": True,
            "auto_merge": False,
        },
        "manual_review_required": True,
    }
    return result


def to_json(result: dict) -> str:
    return json.dumps(result, ensure_ascii=False, indent=2, sort_keys=False) + "\n"


def _md(value: Any) -> str:
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ")


def to_markdown(result: dict) -> str:
    summary = result.get("summary", {})
    lines = [
        f"# 投标审查器 · 多投标文件离线风险预警",
        "",
        f"> 项目：{_md(result.get('project') or '未填写')}  ",
        f"> 生成时间（UTC）：{_md(result.get('generated_at'))}",
        "",
        "## 结论和边界",
        "",
        f"**本报告仅输出风险预警和人工复核线索，不构成串通投标、违法、资格不合格或投标无效认定。**",
        "",
        "检测结果受资料完整性、文件版本、单位换算、招标文件统一模板和报价口径影响；单一报价接近、文本相似、IP/MAC/作者/联系人相同或本地关联线索，均不得自动作出法律结论。",
        "",
        "## 扫描概况",
        "",
        "| 指标 | 数值 |",
        "|---|---:|",
        f"| 投标人 | {summary.get('bidder_count', 0)} |",
        f"| 文件 | {summary.get('file_count', 0)}（成功解析 {summary.get('parsed_file_count', 0)}） |",
        f"| 已识别报价 | {summary.get('quote_count', 0)} |",
        f"| 清单明细（可比/数据不足） | {summary.get('comparable_line_item_count', 0)} / {summary.get('insufficient_line_item_count', 0)} |",
        f"| 风险信号 | {summary.get('signal_count', 0)} |",
        f"| 不支持文件 | {summary.get('unsupported_file_count', 0)} |",
        f"| 解析错误 | {summary.get('parse_error_count', 0)} |",
        f"| 解析提示 | {summary.get('parse_warning_count', 0)} |",
        f"| 跳过文件 / 扫描问题 | {summary.get('skipped_file_count', 0)} / {summary.get('scan_error_count', 0)} |",
        "",
        "## 投标人和来源文件",
        "",
        "| 投标人 | 文件 | 主报价（元） | 明细数 | 元数据字段 |",
        "|---|---|---:|---:|---|",
    ]
    for bidder in result.get("bidders", []):
        quote = bidder.get("primary_quote") or {}
        lines.append(f"| {_md(bidder['name'])} | {len(bidder.get('files', []))} | "
                     f"{quote.get('value', '未识别')} | {bidder.get('line_item_count', 0)} | "
                     f"{_md(', '.join(sorted(bidder.get('metadata', {})))) or '—'} |")
    lines.extend(["", "## 风险信号（全部人工复核）", ""])
    if not result.get("signals"):
        lines.append("未触发本次启用的模式预警；这不等于确认不存在串通或其他违法情形，仍须按业务流程人工核查。")
    for signal in result.get("signals", []):
        lines.extend([
            f"### {signal.get('id')} · {signal.get('rule_id') or '未编号'} · {signal.get('title')}（{signal.get('level')}）",
            "",
            f"- 范围：{_md(signal.get('scope'))}",
            f"- 说明：{_md(signal.get('description'))}",
            f"- 建议：{_md(signal.get('manual_action'))}",
            "- 证据：",
        ])
        for evidence in signal.get("evidence", [])[:20]:
            compact = "; ".join(f"{k}={v}" for k, v in evidence.items() if k not in {"raw", "text"})
            lines.append(f"  - {_md(compact)}")
        lines.append("")
    if (result.get("unsupported_files") or result.get("parse_errors") or result.get("parse_warnings") or
            result.get("skipped_files") or result.get("scan_errors")):
        lines.extend(["## 未完整纳入的资料", ""])
        for file in result.get("unsupported_files", []):
            lines.append(f"- 未支持：`{_md(file.get('path'))}`（{_md(file.get('extension'))}）")
        for error in result.get("parse_errors", []):
            lines.append(f"- 解析失败：`{_md(error.get('path'))}`：{_md(error.get('error'))}")
        for warning in result.get("parse_warnings", []):
            lines.append(f"- 解析提示：`{_md(warning.get('path'))}`：{_md(warning.get('error'))}")
        for skipped in result.get("skipped_files", []):
            lines.append(f"- 已跳过：`{_md(skipped.get('path'))}`（{_md(skipped.get('reason'))}）")
        for error in result.get("scan_errors", []):
            lines.append(f"- 扫描问题：`{_md(error.get('path'))}`：{_md(error.get('error'))}")
        lines.append("")
    lines.extend(["## 规则依据和适用范围", "", "本工具将下列公开规则作为风险线索设计背景，不把算法信号直接等同于法定推定；具体项目还须核对招标文件、交易平台留痕、原始文件和主管部门程序：", ""])
    for source in result.get("legal_notice", {}).get("sources", []):
        lines.append(f"- [{_md(source['name'])}]({source['url']})")
    lines.extend(["", "- 浙江地方规则仅在其适用范围内作为工程建设项目风险设计参考；政府采购、工程建设和其他采购类型不得混用规则。",
                  "- 本报告不是行政处罚、评标否决或责任认定文书；正式结论由有权主体依据完整证据和适用法律程序作出。", ""])
    return "\n".join(lines)


def _resolve_result_path(value: Any, base: Path | None = None) -> Path:
    try:
        path = Path(value).expanduser()
        if base is not None and not path.is_absolute():
            path = base / path
        return path.resolve()
    except (OSError, RuntimeError, TypeError, ValueError) as exc:
        raise ValueError(f"结果中的路径无法解析：{value}") from exc


def _same_file(left: Path, right: Path) -> bool:
    try:
        return os.path.samefile(left, right)
    except (FileNotFoundError, OSError):
        return False


def write_outputs(result: dict, *, json_path: str | Path | None = None, report_path: str | Path | None = None) -> tuple[Path | None, Path | None]:
    """写出用户指定的 JSON/Markdown；不覆盖输入文件，父目录不存在时创建。"""
    input_dir = result.get("input_dir")
    if not input_dir:
        raise ValueError("结果缺少 input_dir，拒绝写入输出")
    root = _resolve_result_path(input_dir)
    source_paths: set[Path] = set()
    for bidder in result.get("bidders", []):
        for source_file in bidder.get("files", []):
            source = source_file.get("path")
            if source:
                source_paths.add(_resolve_result_path(source, root))
    relation_source = result.get("relation_source")
    if isinstance(relation_source, dict) and relation_source.get("path"):
        source_paths.add(_resolve_result_path(relation_source["path"], root))
    for relation in (result.get("relations"), result.get("relation_path")):
        if relation:
            source_paths.add(_resolve_result_path(relation, root))

    pending: list[tuple[Path, str] | None] = []
    for path, content in ((json_path, to_json(result)), (report_path, to_markdown(result))):
        if path is None:
            pending.append(None)
            continue
        target = _resolve_result_path(path)
        if target == root or root in target.parents:
            raise ValueError(f"输出路径不能位于输入目录内：{target}")
        if target in source_paths or any(_same_file(target, source) for source in source_paths):
            raise ValueError(f"输出路径不能覆盖输入源文件：{target}")
        pending.append((target, content))

    outputs: list[Path | None] = []
    for item in pending:
        if item is None:
            outputs.append(None)
            continue
        target, content = item
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(content, encoding="utf-8", newline="\n")
        outputs.append(target)
    return outputs[0], outputs[1]


__all__ = ["SUPPORTED_EXTENSIONS", "THRESHOLDS", "review_directory", "to_json", "to_markdown", "write_outputs"]
